import io
from fastapi import APIRouter, Depends, HTTPException, status
from fastapi.responses import StreamingResponse
from sqlmodel import Session, SQLModel, select
from sqlalchemy import exc
from typing import List, Optional
from io import BytesIO

import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment

from reportlab.lib import colors
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib.units import mm
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle

from app.db.session import get_session
from app.db.models import Chip, Local, Usuario, DetalleVentaChip, IngresoLoteChip
from app.api.modelscreate import ChipCreate
from app.api.modelsupdate import ChipUpdate
from app.api.deps import get_current_user, require_admin, UsuarioActual
from app.api.funciones.fechas import TZ_AR


# ─── SCHEMAS INGRESO LOTE ────────────────────────────────────────────────────

COMPANIAS_VALIDAS = {"CLARO", "PERSONAL", "MOVISTAR", "TUENTI"}


class ChipLoteItem(SQLModel):
    compania: str
    numero_serie: str
    precio: int


class IngresoLoteChipCreate(SQLModel):
    local_id: int
    receptor_id: Optional[int] = None
    observaciones: Optional[str] = None
    chips: List[ChipLoteItem]


router = APIRouter(
    prefix="/chips",
    tags=["Chips"],
)


# ─── HELPERS ──────────────────────────────────────────────────────────────────

def _aplicar_filtros(query, local_id, estado, compania, buscar):
    """Aplica los filtros comunes a una query de Chip."""
    if local_id:
        query = query.where(Chip.local_id == local_id)
    if estado:
        query = query.where(Chip.estado == estado)
    if compania:
        query = query.where(Chip.compania.ilike(f"%{compania}%"))  # type: ignore
    if buscar:
        query = query.where(Chip.numero_serie.ilike(f"%{buscar}%"))  # type: ignore
    return query


# ─── ENDPOINTS ────────────────────────────────────────────────────────────────

@router.get("/export")
def exportar_chips(
    session: Session = Depends(get_session),
    local_id: Optional[int] = None,
    estado: Optional[str] = None,
    compania: Optional[str] = None,
    buscar: Optional[str] = None,
    limit: int = 10000,
    current_user: UsuarioActual = Depends(get_current_user)
):
    """Exportar chips a Excel respetando los filtros activos."""

    # Traer chips y locales en un solo join
    statement = (
        select(Chip, Local)
        .join(Local, Chip.local_id == Local.local_id)  # type: ignore
    )
    statement = _aplicar_filtros(statement, local_id, estado, compania, buscar)
    rows = session.exec(statement.limit(limit)).all()

    # ── Construir el libro ────────────────────────────────────────────────────
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Chips" #type: ignore

    # Estilos de cabecera
    header_font    = Font(bold=True, color="FFFFFF")
    header_fill    = PatternFill("solid", fgColor="1A1A2E")
    header_align   = Alignment(horizontal="center", vertical="center")

    headers = ["ID", "N° de serie", "Compañía", "Precio", "Local", "Estado"]
    for col_idx, h in enumerate(headers, start=1):
        cell = ws.cell(row=1, column=col_idx, value=h) #type: ignore
        cell.font    = header_font
        cell.fill    = header_fill
        cell.alignment = header_align

    # Filas de datos
    for row_idx, (chip, local) in enumerate(rows, start=2):
        ws.cell(row=row_idx, column=1, value=chip.chip_id) #type: ignore
        ws.cell(row=row_idx, column=2, value=chip.numero_serie) #type: ignore
        ws.cell(row=row_idx, column=3, value=chip.compania) #type: ignore
        ws.cell(row=row_idx, column=4, value=chip.precio) #type: ignore
        ws.cell(row=row_idx, column=5, value=local.nombre) #type: ignore
        ws.cell(row=row_idx, column=6, value=chip.estado) #type: ignore

    # Ancho de columnas
    anchos = [8, 22, 16, 12, 20, 14]
    for col_idx, ancho in enumerate(anchos, start=1):
        ws.column_dimensions[ws.cell(row=1, column=col_idx).column_letter].width = ancho #type: ignore

    # ── Serializar y devolver ─────────────────────────────────────────────────
    buffer = BytesIO()
    wb.save(buffer)
    buffer.seek(0)

    return StreamingResponse(
        buffer,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": "attachment; filename=chips.xlsx"},
    )


@router.get("/", response_model=List[Chip])
def listar_chips(
    session: Session = Depends(get_session),
    local_id: Optional[int] = None,
    estado: Optional[str] = None,
    compania: Optional[str] = None,
    buscar: Optional[str] = None,
    skip: int = 0,
    limit: int = 100,
    current_user: UsuarioActual = Depends(get_current_user)
):
    """Listar chips con filtros opcionales. buscar filtra por número de serie."""
    query = select(Chip)
    query = _aplicar_filtros(query, local_id, estado, compania, buscar)
    return session.exec(query.offset(skip).limit(limit)).all()


@router.get("/{chip_id}")
def obtener_chip(
    chip_id: int,
    session: Session = Depends(get_session),
    current_user: UsuarioActual = Depends(get_current_user)
):
    """Obtener un chip específico por ID."""
    chip = session.get(Chip, chip_id)
    if not chip:
        raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail="Chip no encontrado")
    return chip

# deprecado
@router.get("/numero-serie/{numero_serie}")
def buscar_por_numero_serie(
    numero_serie: str,
    session: Session = Depends(get_session),
    current_user: UsuarioActual = Depends(get_current_user)
):
    """Buscar chip por número de serie exacto."""
    chip = session.exec(
        select(Chip).where(Chip.numero_serie == numero_serie)
    ).first()
    if not chip:
        raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail="Chip no encontrado")
    return chip


@router.post("/", status_code=status.HTTP_201_CREATED)
def crear_chip(
    chip_in: ChipCreate,
    session: Session = Depends(get_session),
    current_user: UsuarioActual = Depends(require_admin)
):
    """Crear un nuevo chip."""
    chip = Chip(**chip_in.model_dump())
    session.add(chip)

    try:
        session.commit()
        session.refresh(chip)
        return chip

    except exc.IntegrityError as e:
        session.rollback()
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail=str(e)
        )


@router.put("/{chip_id}")
def actualizar_chip(
    chip_id: int,
    chip_in: ChipUpdate,
    session: Session = Depends(get_session),
    current_user: UsuarioActual = Depends(require_admin)
):
    """Actualizar un chip."""
    chip = session.get(Chip, chip_id)
    if not chip:
        raise HTTPException(
            status_code=status.HTTP_404_NOT_FOUND,
            detail="Chip no encontrado"
        )

    for field, value in chip_in.model_dump(exclude_unset=True).items():
        setattr(chip, field, value)

    session.add(chip)

    try:
        session.commit()
        session.refresh(chip)
        return chip

    except exc.IntegrityError as e:
        session.rollback()
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail=str(e)
        )


@router.delete("/{chip_id}")
def eliminar_chip(
    chip_id: int,
    session: Session = Depends(get_session),
    current_user: UsuarioActual = Depends(require_admin)
):
    """
    Elimina un chip físicamente.
    Si el chip tiene ventas asociadas (detalles_ventas_chips),
    la operación es rechazada — en ese caso actualizá el estado a 'VENDIDO' en su lugar.
    """
    chip = session.get(Chip, chip_id)
    if not chip:
        raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail="Chip no encontrado")

    try:
        session.delete(chip)
        session.commit()
        return {"ok": True, "mensaje": "Chip eliminado exitosamente"}
    except exc.IntegrityError:
        session.rollback()
        raise HTTPException(
            status_code=status.HTTP_409_CONFLICT,
            detail=(
                "No se puede eliminar el chip porque tiene ventas asociadas. "
                "Actualizá el estado a 'VENDIDO' en su lugar."
            )
        )
    except Exception as e:
        session.rollback()
        raise HTTPException(
            status_code=status.HTTP_500_INTERNAL_SERVER_ERROR,
            detail=f"Error inesperado: {str(e)}"
        )


# ─── INGRESO POR LOTE ────────────────────────────────────────────────────────

@router.get("/ingresos/")
def listar_ingresos_chips(
    session: Session = Depends(get_session),
    local_id: Optional[int] = None,
    fecha_desde: Optional[str] = None,
    fecha_hasta: Optional[str] = None,
    skip: int = 0,
    limit: int = 20,
    current_user: UsuarioActual = Depends(require_admin),
):
    """Lista los ingresos de lotes de chips con filtros opcionales."""
    query = (
        select(IngresoLoteChip, Local, Usuario)
        .join(Local, IngresoLoteChip.local_id == Local.local_id, isouter=True)  # type: ignore
        .join(Usuario, IngresoLoteChip.receptor_id == Usuario.usuario_id, isouter=True)  # type: ignore
        .order_by(IngresoLoteChip.ingreso_lote_chip_id.desc())  # type: ignore
    )
    if local_id:
        query = query.where(IngresoLoteChip.local_id == local_id)
    if fecha_desde:
        query = query.where(IngresoLoteChip.fecha >= fecha_desde) #type: ignore
    if fecha_hasta:
        query = query.where(IngresoLoteChip.fecha <= fecha_hasta + " 23:59:59") #type: ignore
    rows = session.exec(query.offset(skip).limit(limit)).all()
    return [
        {
            "ingreso_lote_chip_id": lote.ingreso_lote_chip_id,
            "fecha": lote.fecha,
            "local": local.nombre if local else None,
            "nombre_receptor": receptor.nombre if receptor else None,
            "observaciones": lote.observaciones,
        }
        for lote, local, receptor in rows
    ]


@router.get("/ingresos/{ingreso_lote_chip_id}")
def detalle_ingreso_chips(
    ingreso_lote_chip_id: int,
    session: Session = Depends(get_session),
    current_user: UsuarioActual = Depends(require_admin),
):
    """Devuelve el detalle de un ingreso de lote de chips con los chips incluidos."""
    lote = session.get(IngresoLoteChip, ingreso_lote_chip_id)
    if not lote:
        raise HTTPException(status_code=404, detail="Ingreso no encontrado.")
    local    = session.get(Local,    lote.local_id)    if lote.local_id    else None
    receptor = session.get(Usuario,  lote.receptor_id) if lote.receptor_id else None
    chips = session.exec(
        select(Chip).where(Chip.ingreso_lote_chip_id == ingreso_lote_chip_id)
    ).all()
    return {
        "ingreso_lote_chip_id": lote.ingreso_lote_chip_id,
        "fecha": lote.fecha,
        "local": local.nombre if local else None,
        "nombre_receptor": receptor.nombre if receptor else None,
        "observaciones": lote.observaciones,
        "chips": [
            {"numero_serie": c.numero_serie, "compania": c.compania, "precio": c.precio}
            for c in chips
        ],
    }


@router.post("/ingresar-lote", status_code=status.HTTP_201_CREATED)
def ingresar_lote_chips(
    payload: IngresoLoteChipCreate,
    session: Session = Depends(get_session),
    current_user: UsuarioActual = Depends(require_admin),
):
    """
    Crea múltiples chips de una vez y los vincula a un IngresoLoteChip.
    Devuelve el ingreso_lote_chip_id para poder descargar el remito PDF.
    """
    if not payload.chips:
        raise HTTPException(status_code=status.HTTP_400_BAD_REQUEST,
                            detail="Debe incluir al menos un chip.")

    companias_invalidas = [c.compania for c in payload.chips if c.compania.upper() not in COMPANIAS_VALIDAS]
    if companias_invalidas:
        raise HTTPException(status_code=status.HTTP_400_BAD_REQUEST,
                            detail=f"Compañía inválida: {', '.join(set(companias_invalidas))}. Valores permitidos: {', '.join(sorted(COMPANIAS_VALIDAS))}.")

    # Verificar duplicados dentro del payload
    series = [c.numero_serie.strip() for c in payload.chips]
    if len(series) != len(set(series)):
        raise HTTPException(status_code=status.HTTP_400_BAD_REQUEST,
                            detail="Hay números de serie duplicados en el lote.")

    # Verificar local
    local = session.get(Local, payload.local_id)
    if not local:
        raise HTTPException(status_code=status.HTTP_404_NOT_FOUND,
                            detail=f"Local {payload.local_id} no encontrado.")

    try:
        lote = IngresoLoteChip(
            local_id=payload.local_id,
            receptor_id=payload.receptor_id,
            usuario_id=current_user.usuario_id,
            observaciones=payload.observaciones or None,
        )
        session.add(lote)
        session.flush()  # obtener ingreso_lote_chip_id

        chips_creados = []
        for item in payload.chips:
            # Verificar que el número de serie no exista ya en la BD
            existente = session.exec(
                select(Chip).where(Chip.numero_serie == item.numero_serie.strip())
            ).first()
            if existente:
                raise HTTPException(
                    status_code=status.HTTP_400_BAD_REQUEST,
                    detail=f"El número de serie '{item.numero_serie}' ya existe en el sistema."
                )
            chip = Chip(
                compania=item.compania.strip(),
                numero_serie=item.numero_serie.strip(),
                precio=item.precio,
                local_id=payload.local_id,
                estado="DISPONIBLE",
                ingreso_lote_chip_id=lote.ingreso_lote_chip_id,
            )
            session.add(chip)
            chips_creados.append(item)

        session.commit()
        session.refresh(lote)

        return {
            "mensaje": "Ingreso de lote de chips realizado exitosamente.",
            "ingreso_lote_chip_id": lote.ingreso_lote_chip_id,
            "fecha": lote.fecha,
            "local": local.nombre,
            "total_chips": len(chips_creados),
        }

    except HTTPException:
        session.rollback()
        raise
    except exc.IntegrityError:
        session.rollback()
        raise HTTPException(status_code=status.HTTP_400_BAD_REQUEST,
                            detail="Error de integridad: algún número de serie ya existe.")
    except Exception as e:
        session.rollback()
        raise HTTPException(status_code=status.HTTP_500_INTERNAL_SERVER_ERROR,
                            detail=f"Error inesperado: {str(e)}")


@router.get("/ingresos/{ingreso_lote_chip_id}/pdf")
def remito_ingreso_chips_pdf(
    ingreso_lote_chip_id: int,
    session: Session = Depends(get_session),
    current_user: UsuarioActual = Depends(require_admin),
):
    """Genera el remito PDF de un ingreso de lote de chips."""
    lote = session.get(IngresoLoteChip, ingreso_lote_chip_id)
    if not lote:
        raise HTTPException(status_code=404, detail="Ingreso no encontrado.")

    local    = session.get(Local,    lote.local_id)    if lote.local_id    else None
    receptor = session.get(Usuario,  lote.receptor_id) if lote.receptor_id else None

    chips = session.exec(
        select(Chip).where(Chip.ingreso_lote_chip_id == ingreso_lote_chip_id)
    ).all()

    total_chips = len(chips)
    fecha_str   = lote.fecha.astimezone(TZ_AR).strftime("%d/%m/%Y %H:%M") if lote.fecha else "-"
    local_str   = local.nombre    if local    else "-"
    receptor_str = receptor.nombre if receptor else "-"

    # ── PDF ────────────────────────────────────────────────────────────────────
    buffer = io.BytesIO()
    W = A4[0] - 30 * mm
    doc = SimpleDocTemplate(
        buffer, pagesize=A4,
        leftMargin=15*mm, rightMargin=15*mm,
        topMargin=12*mm, bottomMargin=12*mm,
    )

    styles = getSampleStyleSheet()
    normal = ParagraphStyle("n",  parent=styles["Normal"], fontSize=9,  leading=11)
    bold   = ParagraphStyle("b",  parent=normal, fontName="Helvetica-Bold")
    small  = ParagraphStyle("sm", parent=normal, fontSize=8)
    title  = ParagraphStyle("t",  parent=styles["Normal"], fontSize=13,
                             fontName="Helvetica-Bold", spaceAfter=6)
    hcell  = ParagraphStyle("hc", parent=bold, fontSize=8.5)
    cell   = ParagraphStyle("c",  parent=normal, fontSize=8.5, leading=10)

    story = []

    # Encabezado
    story.append(Paragraph(f"REMITO DE INGRESO DE CHIPS  #{ingreso_lote_chip_id}", title))

    meta = Table(
        [
            [Paragraph("<b>Fecha:</b>",    bold), Paragraph(fecha_str,    normal),
             Paragraph("<b>Local:</b>",    bold), Paragraph(local_str,    normal)],
            [Paragraph("<b>Receptor:</b>", bold), Paragraph(receptor_str, normal),
             Paragraph("",                bold), Paragraph("",            normal)],
        ],
        colWidths=[W*0.14, W*0.36, W*0.14, W*0.36],
    )
    meta.setStyle(TableStyle([
        ("FONTSIZE",      (0,0),(-1,-1), 9),
        ("TOPPADDING",    (0,0),(-1,-1), 2),
        ("BOTTOMPADDING", (0,0),(-1,-1), 2),
        ("LEFTPADDING",   (0,0),(-1,-1), 3),
        ("RIGHTPADDING",  (0,0),(-1,-1), 3),
        ("BOX",           (0,0),(-1,-1), 0.5, colors.black),
        ("LINEBELOW",     (0,0),(-1,0),  0.5, colors.black),
        ("LINEBEFORE",    (2,0),(2,-1),  0.5, colors.black),
    ]))
    story.append(meta)
    story.append(Spacer(1, 10*mm))

    # Tabla de chips
    header = [Paragraph(t, hcell) for t in ["N° de serie", "Compañía", "Precio"]]
    data   = [header]
    for chip in chips:
        data.append([
            Paragraph(chip.numero_serie, cell),
            Paragraph(chip.compania,     cell),
            Paragraph(f"${chip.precio:,}".replace(",", "."), cell),
        ])

    tbl = Table(data, colWidths=[W*0.45, W*0.30, W*0.25], repeatRows=1)
    tbl.setStyle(TableStyle([
        ("FONTSIZE",      (0,0),(-1,-1), 8.5),
        ("GRID",          (0,0),(-1,-1), 0.25, colors.black),
        ("LINEBELOW",     (0,0),(-1,0),  0.75, colors.black),
        ("FONTNAME",      (0,0),(-1,0),  "Helvetica-Bold"),
        ("VALIGN",        (0,0),(-1,-1), "TOP"),
        ("LEFTPADDING",   (0,0),(-1,-1), 3),
        ("RIGHTPADDING",  (0,0),(-1,-1), 3),
        ("TOPPADDING",    (0,0),(-1,-1), 2),
        ("BOTTOMPADDING", (0,0),(-1,-1), 2),
        ("ALIGN",         (2,0),(2,-1),  "RIGHT"),
    ]))
    story.append(tbl)
    story.append(Spacer(1, 4*mm))

    # Total
    tot = Table(
        [[Paragraph(f"Total chips ingresados: <b>{total_chips}</b>", small)]],
        colWidths=[W],
    )
    tot.setStyle(TableStyle([
        ("ALIGN",         (0,0),(-1,-1), "RIGHT"),
        ("TOPPADDING",    (0,0),(-1,-1), 2),
        ("BOTTOMPADDING", (0,0),(-1,-1), 2),
        ("LEFTPADDING",   (0,0),(-1,-1), 3),
        ("RIGHTPADDING",  (0,0),(-1,-1), 3),
    ]))
    story.append(tot)
    story.append(Spacer(1, 14*mm))

    # Firma
    firma = Table(
        [
            [Paragraph("Firma y aclaración:", bold),
             Paragraph("_" * 40, normal)],
            [Paragraph("", normal), Paragraph("", small)],
        ],
        colWidths=[W*0.20, W*0.80],
    )
    firma.setStyle(TableStyle([
        ("TOPPADDING",    (0,0),(-1,-1), 1),
        ("BOTTOMPADDING", (0,0),(-1,-1), 1),
        ("LEFTPADDING",   (0,0),(-1,-1), 0),
        ("FONTSIZE",      (0,0),(-1,-1), 9),
    ]))
    story.append(firma)

    if lote.observaciones:
        story.append(Spacer(1, 6*mm))
        story.append(Paragraph(f"<b>Observaciones:</b> {lote.observaciones}", small))

    doc.build(story)
    filename = f"ingreso_chips_{ingreso_lote_chip_id}_{fecha_str.replace('/', '-').replace(' ', '_').replace(':', '')}.pdf"
    return StreamingResponse(
        io.BytesIO(buffer.getvalue()),
        media_type="application/pdf",
        headers={"Content-Disposition": f'attachment; filename="{filename}"'},
    )