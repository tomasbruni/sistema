#PARA OPERACIONES MANUALES DE STOCK
from fastapi import APIRouter, Depends, HTTPException, status
from sqlmodel import Session

from typing import Optional
from sqlmodel import Session, SQLModel, select, col, Field
from sqlalchemy import exc

from app.db.session import get_session, SessionDep
from app.db.models import *

from app.api.modelscreate import *
from app.api.modelsupdate import *

from fastapi.responses import StreamingResponse
from io import BytesIO
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment
from datetime import date, datetime

from app.api.deps import get_current_user, require_admin, UsuarioActual
from app.api.funciones.movimientos_stock import aplicar_movimiento_stock
import time


router = APIRouter(
    prefix="/stock",
    tags=["OPERACIONES MANUALES DE STOCK"],
)


class StockResponse(SQLModel):
    """Respuesta con información de stock."""
    id: int
    accesorio_id: int
    local_id: int
    cantidad: int
    accesorio_nombre: Optional[str] = None
    local_nombre: Optional[str] = None

@router.get("/export")
def exportar_stock(
    buscar: Optional[str] = "",
    tipo_id: Optional[int] = None,
    subtipo_id: Optional[int] = None,
    local_id: Optional[int] = None,
    activo: Optional[bool] = True,
    mostrar_listado: bool = True,
    session: Session = Depends(get_session),
    current_user: UsuarioActual = Depends(require_admin)
):
    statement = (
        select(StockAccesorio, Accesorio, Local)
        .join(Accesorio, StockAccesorio.accesorio_id == Accesorio.accesorio_id) #type: ignore
        .join(Local, StockAccesorio.local_id == Local.local_id) #type: ignore
    )
    if buscar:
        statement = statement.where(Accesorio.nombre.ilike(f"%{buscar}%")) #type: ignore
    if tipo_id is not None:
        statement = statement.where(Accesorio.tipo_id == tipo_id)
    if subtipo_id is not None:
        statement = statement.where(Accesorio.subtipo_id == subtipo_id)
    if local_id is not None:
        statement = statement.where(StockAccesorio.local_id == local_id)
    if activo is not None:
        statement = statement.where(Accesorio.activo == activo)
    resultados = session.exec(statement).all()

    # ── Nombres de filtros ────────────────────────────────────────────────────
    local_nombre = "Todos"
    if local_id is not None:
        local_obj = session.get(Local, local_id)
        if local_obj:
            local_nombre = local_obj.nombre

    tipo_nombre = None
    if tipo_id is not None:
        tipo_obj = session.get(TipoAccesorio, tipo_id)
        tipo_nombre = tipo_obj.nombre if tipo_obj else str(tipo_id)

    subtipo_nombre = None
    if subtipo_id is not None:
        subtipo_obj = session.get(SubtipoAccesorio, subtipo_id)
        subtipo_nombre = subtipo_obj.nombre if subtipo_obj else str(subtipo_id)

    # ── Totales por subtipo (solo cuando hay tipo pero no subtipo) ────────────
    subtipo_totales: dict = {}
    if tipo_id is not None and subtipo_id is None:
        subtipos = session.exec(
            select(SubtipoAccesorio).where(SubtipoAccesorio.tipo_id == tipo_id)
        ).all()
        subtipo_map: dict = {s.subtipo_id: s.nombre for s in subtipos}
        subtipo_map[None] = "(Sin subtipo)"
        for stock, accesorio, _local in resultados:
            sid = accesorio.subtipo_id
            if sid not in subtipo_totales:
                subtipo_totales[sid] = {"nombre": subtipo_map.get(sid, str(sid)), "cantidad": 0}
            subtipo_totales[sid]["cantidad"] += stock.cantidad

    total_cantidad = sum(stock.cantidad for stock, _, _ in resultados)

    # ── Workbook ──────────────────────────────────────────────────────────────
    hoy = date.today().strftime("%Y%m%d")
    local_fn = local_nombre.lower().replace(" ", "_") if local_id else "todos"
    filename = f"stock_{local_fn}_{hoy}.xlsx"

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Stock" #type: ignore

    DARK = "1A1A2E"
    MID  = "3A3A5E"

    def hdr(row, col, value):
        c = ws.cell(row=row, column=col, value=value) #type: ignore
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill(fill_type="solid", fgColor=DARK)
        c.alignment = Alignment(horizontal="center")
        return c

    def lbl(row, col, value):
        c = ws.cell(row=row, column=col, value=value) #type: ignore
        c.font = Font(bold=True)
        return c

    def summary_hdr(row, col, value):
        c = ws.cell(row=row, column=col, value=value) #type: ignore
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill(fill_type="solid", fgColor=MID)
        c.alignment = Alignment(horizontal="center")
        return c

    r = 1

    # Título
    ws.merge_cells(f"A{r}:E{r}") #type: ignore
    c = ws.cell(row=r, column=1, value="REPORTE DE STOCK") #type: ignore
    c.font = Font(bold=True, size=14, color="FFFFFF")
    c.fill = PatternFill(fill_type="solid", fgColor=DARK)
    c.alignment = Alignment(horizontal="center")
    ws.row_dimensions[r].height = 22 #type: ignore
    r += 1

    # Metadatos
    lbl(r, 1, "Fecha de generación:"); ws.cell(row=r, column=2, value=datetime.now().strftime("%d/%m/%Y %H:%M")); r += 1 #type: ignore
    lbl(r, 1, "Local:");               ws.cell(row=r, column=2, value=local_nombre); r += 1 #type: ignore
    lbl(r, 1, "Tipo:");                ws.cell(row=r, column=2, value=tipo_nombre or "Todos"); r += 1 #type: ignore
    if tipo_id is not None:
        lbl(r, 1, "Subtipo:"); ws.cell(row=r, column=2, value=subtipo_nombre or "Todos"); r += 1 #type: ignore

    r += 1  # fila vacía

    # Resumen
    if subtipo_totales:
        summary_hdr(r, 1, "Subtipo"); summary_hdr(r, 2, "Cantidad"); r += 1
        for sid, data in sorted(subtipo_totales.items(), key=lambda x: (x[0] is None, x[1]["nombre"])):
            ws.cell(row=r, column=1, value=data["nombre"]); ws.cell(row=r, column=2, value=data["cantidad"]); r += 1 #type: ignore
        lbl(r, 1, "TOTAL"); c = ws.cell(row=r, column=2, value=total_cantidad); c.font = Font(bold=True); r += 1 #type: ignore
    else:
        lbl(r, 1, "Total accesorios:"); ws.cell(row=r, column=2, value=total_cantidad); r += 1 #type: ignore

    # Listado detallado
    if mostrar_listado:
        r += 1  # fila vacía
        for col_idx, titulo in enumerate(["ID", "Accesorio", "Cantidad", "Local"], start=1):
            hdr(r, col_idx, titulo)
        r += 1
        for stock, accesorio, local in resultados:
            ws.cell(row=r, column=1, value=accesorio.accesorio_id) #type: ignore
            ws.cell(row=r, column=2, value=accesorio.nombre) #type: ignore
            ws.cell(row=r, column=3, value=stock.cantidad) #type: ignore
            ws.cell(row=r, column=4, value=local.nombre) #type: ignore
            r += 1

    anchos = {"A": 28, "B": 40, "C": 20, "D": 12}
    for col_letra, ancho in anchos.items():
        ws.column_dimensions[col_letra].width = ancho #type: ignore

    buffer = BytesIO()
    wb.save(buffer)
    buffer.seek(0)

    return StreamingResponse(
        buffer,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="{filename}"'}
    )


@router.get("/export-por-exclusion")
def exportar_stock_por_exclusion(
    excluir_tipo_ids: Optional[str] = None,
    local_id: Optional[int] = None,
    session: Session = Depends(get_session),
    current_user: UsuarioActual = Depends(require_admin)
):
    """Exporta todo el stock activo excluyendo los tipos indicados."""
    excluir_ids: list[int] = []
    if excluir_tipo_ids:
        try:
            excluir_ids = [int(x) for x in excluir_tipo_ids.split(",") if x.strip()]
        except ValueError:
            pass

    statement = (
        select(StockAccesorio, Accesorio, Local)
        .join(Accesorio, StockAccesorio.accesorio_id == Accesorio.accesorio_id)  # type: ignore
        .join(Local, StockAccesorio.local_id == Local.local_id)  # type: ignore
        .where(Accesorio.activo == True)  # type: ignore
    )

    if local_id is not None:
        statement = statement.where(StockAccesorio.local_id == local_id)
    if excluir_ids:
        statement = statement.where(Accesorio.tipo_id.notin_(excluir_ids))  # type: ignore

    resultados = session.exec(statement).all()

    nombre_local = "todos"
    if local_id is not None:
        local_obj = session.get(Local, local_id)
        if local_obj:
            nombre_local = local_obj.nombre.lower().replace(" ", "_")

    hoy = date.today().strftime("%Y%m%d")
    filename = f"stock_{nombre_local}_{hoy}.xlsx"

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Stock"  # type: ignore

    header_font  = Font(bold=True, color="FFFFFF")
    header_fill  = PatternFill(fill_type="solid", fgColor="1A1A2E")
    header_align = Alignment(horizontal="center")

    columnas = ["ID", "Accesorio", "Cantidad", "Local"]
    for col_idx, titulo in enumerate(columnas, start=1):
        cell = ws.cell(row=1, column=col_idx, value=titulo)  # type: ignore
        cell.font      = header_font
        cell.fill      = header_fill
        cell.alignment = header_align

    for stock, accesorio, local in resultados:
        ws.append([  # type: ignore
            accesorio.accesorio_id,
            accesorio.nombre,
            stock.cantidad,
            local.nombre,
        ])

    anchos = {"A": 14, "B": 40, "C": 20, "D": 12}
    for col_letra, ancho in anchos.items():
        ws.column_dimensions[col_letra].width = ancho  # type: ignore

    buffer = BytesIO()
    wb.save(buffer)
    buffer.seek(0)

    return StreamingResponse(
        buffer,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="{filename}"'}
    )


@router.get("/", response_model=list[StockResponse])
def listar_stock(
    skip: int = 0,
    limit: int = 100,
    buscar: Optional[str] = "",
    tipo_id: Optional[int] = None,
    subtipo_id: Optional[int] = None,
    local_id: Optional[int] = None,
    activo: Optional[bool] = True,
    session: Session = Depends(get_session),
    current_user: UsuarioActual = Depends(get_current_user)
):
    statement = (
        select(StockAccesorio, Accesorio, Local)
        .join(Accesorio, StockAccesorio.accesorio_id == Accesorio.accesorio_id) #type: ignore
        .join(Local, StockAccesorio.local_id == Local.local_id).where(Local.activo == True) #type: ignore
    )

    if buscar:
        statement = statement.where(Accesorio.nombre.ilike(f"%{buscar}%")) #type: ignore
    if tipo_id is not None:
        statement = statement.where(Accesorio.tipo_id == tipo_id)
    if subtipo_id is not None:
        statement = statement.where(Accesorio.subtipo_id == subtipo_id)
    if local_id is not None:
        statement = statement.where(StockAccesorio.local_id == local_id)
    if activo is not None:
        statement = statement.where(Accesorio.activo == activo)

    statement = statement.offset(skip).limit(limit)
    resultados = session.exec(statement).all()

    response = [
        StockResponse(
            id=stock.stock_id, #type: ignore
            accesorio_id=stock.accesorio_id,
            local_id=stock.local_id,
            cantidad=stock.cantidad,
            accesorio_nombre=accesorio.nombre,
            local_nombre=local.nombre
        )
        for stock, accesorio, local in resultados
    ]
    return response


@router.get("/{stock_id}", response_model=StockResponse)
def obtener_stock(
    stock_id: int,
    session: Session = Depends(get_session),
    current_user: UsuarioActual = Depends(get_current_user)
):
    stock = session.get(StockAccesorio, stock_id)
    if not stock:
        raise HTTPException(
            status_code=status.HTTP_404_NOT_FOUND,
            detail=f"Stock con ID {stock_id} no encontrado"
        )
    accesorio = session.get(Accesorio, stock.accesorio_id)
    local = session.get(Local, stock.local_id)

    if accesorio and local:
        return StockResponse(
            id=stock.stock_id, #type: ignore
            accesorio_id=stock.accesorio_id,
            local_id=stock.local_id,
            cantidad=stock.cantidad,
            accesorio_nombre=accesorio.nombre,
            local_nombre=local.nombre
        )
    else:
        raise HTTPException(
            status_code=status.HTTP_404_NOT_FOUND,
            detail="No tiene accesorio asociado"
        )


@router.post("/ajustar", status_code=status.HTTP_200_OK)
def ajustar_stock(
    ajuste: AjusteStockCreate,
    session: Session = Depends(get_session),
    current_user: UsuarioActual = Depends(require_admin)
):
    """
    Ajustar stock a una cantidad específica (por ej: después de inventario físico).
    Crea un movimiento de tipo AJUSTE con la diferencia entre el stock actual y la nueva cantidad.
    """
    if ajuste.cantidad_nueva < 0:
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail="La cantidad debe ser positiva"
        )

    try:
        stock_query = select(StockAccesorio).where(
            StockAccesorio.accesorio_id == ajuste.accesorio_id,
            StockAccesorio.local_id == ajuste.local_id
        ).with_for_update()
        
        stock = session.exec(stock_query).first()

        # time.sleep(20) # testing

        if not stock:
            raise HTTPException(
                status_code=status.HTTP_404_NOT_FOUND,
                detail="Stock no encontrado para este accesorio y local"
            )

        cantidad_original = stock.cantidad
        diferencia = ajuste.cantidad_nueva - cantidad_original

        movimiento = aplicar_movimiento_stock(
            session, stock,
            tipo_movimiento=TipoMovimiento.AJUSTE,
            cantidad=diferencia, #cambio en la logica, bien
            motivo=f"{ajuste.motivo} (Ajuste: {cantidad_original} → {ajuste.cantidad_nueva})",
            usuario_id=current_user.usuario_id,
        )
        session.commit()
        session.refresh(movimiento)

        return {
            "mensaje": "Ajuste realizado exitosamente",
            "stock_anterior": cantidad_original,
            "stock_nuevo": ajuste.cantidad_nueva,
            "diferencia": diferencia,
            "movimiento_id": movimiento.id,
            "fecha": movimiento.fecha,
        }

    except exc.IntegrityError as e:
        session.rollback()
        raise HTTPException(status_code=status.HTTP_400_BAD_REQUEST, detail=f"Database error: {e.orig}")
    except HTTPException:
        session.rollback()
        raise
    except Exception as e:
        session.rollback()
        raise HTTPException(status_code=status.HTTP_500_INTERNAL_SERVER_ERROR, detail=f"Error inesperado: {str(e)}")


# ─── EGRESO INDIVIDUAL (se mantiene para salidas manuales unitarias) ──────────
@router.post("/egreso", status_code=status.HTTP_200_OK)
def egreso_stock(
    movimiento: MovimientoStockCreate,
    session: Session = Depends(get_session),
    current_user: UsuarioActual = Depends(require_admin)
):
    """
    Egreso manual individual de stock.
    Solo acepta tipo_movimiento = SALIDA.
    """
    if movimiento.tipo_movimiento != TipoMovimiento.SALIDA:
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail="Este endpoint solo acepta movimientos de tipo SALIDA"
        )

    if movimiento.cantidad <= 0:
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail="La cantidad debe ser mayor a cero"
        )

    try:
        stock_query = select(StockAccesorio).where(
            StockAccesorio.accesorio_id == movimiento.accesorio_id,
            StockAccesorio.local_id == movimiento.local_id
        ).with_for_update()

        stock = session.exec(stock_query).first()

        if not stock:
            raise HTTPException(
                status_code=status.HTTP_400_BAD_REQUEST,
                detail=f"No existe stock del accesorio con ID {movimiento.accesorio_id}"
            )

        if stock.cantidad < movimiento.cantidad:
            raise HTTPException(
                status_code=status.HTTP_400_BAD_REQUEST,
                detail=f"Stock insuficiente. Disponible: {stock.cantidad}, Solicitado: {movimiento.cantidad}"
            )

        motivo = f"Egreso manual en local {movimiento.local_id}"
        if movimiento.motivo:
            motivo += f" - {movimiento.motivo}"

        mov_nuevo = aplicar_movimiento_stock(
            session, stock,
            tipo_movimiento=TipoMovimiento.SALIDA,
            cantidad=-abs(movimiento.cantidad),
            motivo=motivo,
            usuario_id=current_user.usuario_id,
        )
        session.commit()
        session.refresh(mov_nuevo)

        return {
            "mensaje": "Egreso realizado exitosamente",
            "movimiento": {
                "id": mov_nuevo.id,
                "tipo": "SALIDA",
                "cantidad": movimiento.cantidad,
                "fecha": mov_nuevo.fecha
            },
            "stock": {
                "local_id": movimiento.local_id,
                "accesorio_id": movimiento.accesorio_id,
                "cantidad_actual": stock.cantidad
            }
        }

    except exc.IntegrityError:
        session.rollback()
        raise HTTPException(status_code=status.HTTP_400_BAD_REQUEST, detail="Error en la integridad de datos")
    except HTTPException:
        session.rollback()
        raise
    except Exception as e:
        session.rollback()
        raise HTTPException(status_code=status.HTTP_500_INTERNAL_SERVER_ERROR, detail=f"Error inesperado: {str(e)}")


# ─── INGRESO POR LOTE ─────────────────────────────────────────────────────────
@router.post("/ingresar-lote", status_code=status.HTTP_200_OK)
def ingresar_lote(
    ingreso_lote: IngresoLoteCreate,
    session: Session = Depends(get_session),
    current_user: UsuarioActual = Depends(require_admin)
):
    """
    Ingresa múltiples accesorios al stock de un local en una sola operación.
    Crea un registro IngresoLote + un MovimientoStock de ENTRADA por cada producto.
    Devuelve el detalle completo para que el frontend genere el PDF.
    """
    if not ingreso_lote.productos:
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail="Debe incluir al menos un producto en el lote"
        )

    # Verificar que el local existe
    local = session.get(Local, ingreso_lote.local_id)
    if not local:
        raise HTTPException(
            status_code=status.HTTP_404_NOT_FOUND,
            detail=f"Local {ingreso_lote.local_id} no encontrado"
        )

    # Verificar duplicados en el payload
    ids = [item.accesorio_id for item in ingreso_lote.productos]
    if len(ids) != len(set(ids)):
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail="No se permiten accesorios duplicados en el lote. Consolidá las cantidades en una sola línea."
        )

    try:
        # Crear registro de lote
        lote = IngresoLote(
            usuario_id=current_user.usuario_id,
            receptor_id=ingreso_lote.receptor_id,
            observaciones=ingreso_lote.observaciones or None,
            local_id=ingreso_lote.local_id,
        )
        session.add(lote)
        session.flush()  # obtenemos ingreso_lote_id

        resultado_items = []
        # aca pasa lo mismo q en ventas, si alguna comprobacion falla, 
        # el id de IngresoLote se incrementa igual
        for item in ingreso_lote.productos:
            if item.cantidad_ingreso <= 0:
                raise HTTPException(
                    status_code=status.HTTP_400_BAD_REQUEST,
                    detail=f"La cantidad del accesorio {item.accesorio_id} debe ser mayor a cero"
                )

            accesorio = session.get(Accesorio, item.accesorio_id)
            if not accesorio:
                raise HTTPException(
                    status_code=status.HTTP_404_NOT_FOUND,
                    detail=f"Accesorio ID {item.accesorio_id} no encontrado"
                )
            if not accesorio.activo:
                raise HTTPException(
                    status_code=status.HTTP_400_BAD_REQUEST,
                    detail=f"El accesorio '{accesorio.nombre}' no está activo"
                )

            # Obtener o crear el stock con lock
            stock_query = select(StockAccesorio).where(
                StockAccesorio.accesorio_id == item.accesorio_id,
                StockAccesorio.local_id == ingreso_lote.local_id
            ).with_for_update()


            stock = session.exec(stock_query).first()
            #si no existe stock (no deberia pasar igual) se crea
            if not stock:
                stock = StockAccesorio(
                    accesorio_id=item.accesorio_id,
                    local_id=ingreso_lote.local_id,
                    cantidad=0
                )
                session.add(stock)
                session.flush()

            if (ingreso_lote.receptor_id == 1):
                time.sleep(20) #para testing

            #creo movimientos de entrada asociados al ingreso
            mov = aplicar_movimiento_stock(
                session, stock,
                tipo_movimiento=TipoMovimiento.ENTRADA,
                cantidad=item.cantidad_ingreso,
                motivo=f"Ingreso lote #{lote.ingreso_lote_id} - {local.nombre}",
                usuario_id=current_user.usuario_id,
                ingreso_lote_id=lote.ingreso_lote_id,
            )

            # esto despues va al front para generar el remito
            # asi no tengo que hacer un endpoint q haga joins
            # igual capaz podria tener uno q genere excels en el back
            # asi puedo generar reportes de un ingreso cuando se quiera
            resultado_items.append({
                "accesorio_id": item.accesorio_id,
                "accesorio_nombre": accesorio.nombre,
                "cantidad_ingresada": item.cantidad_ingreso,
                "stock_anterior": mov.stock_anterior,
                "stock_nuevo": mov.stock_nuevo,
            })

        session.commit()
        session.refresh(lote)

        return {
            "mensaje": "Ingreso de lote realizado exitosamente",
            "ingreso_lote_id": lote.ingreso_lote_id,
            "fecha": lote.fecha,
            "local_id": ingreso_lote.local_id,
            "local_nombre": local.nombre,
            "usuario_id": current_user.usuario_id,
            "observaciones": lote.observaciones,
            "items": resultado_items,
        }

    except exc.IntegrityError as e:
        session.rollback()
        raise HTTPException(status_code=status.HTTP_400_BAD_REQUEST, detail=f"Error de integridad: {e.orig}")
    except HTTPException:
        session.rollback()
        raise
    except Exception as e:
        session.rollback()
        raise HTTPException(status_code=status.HTTP_500_INTERNAL_SERVER_ERROR, detail=f"Error inesperado: {str(e)}")

# siempre se usa usuario_id = current_user para crear los movimientos
# ─── TRANSFERENCIA POR LOTE ───────────────────────────────────────────────────
@router.post("/transferir-lote", status_code=status.HTTP_200_OK)
def transferir_lote(
    transferencia: TransferenciaLoteCreate,
    session: Session = Depends(get_session),
    current_user: UsuarioActual = Depends(get_current_user)
):
    """
    Transfiere múltiples accesorios entre dos locales en una sola operación.
    Crea un registro Transferencia + dos MovimientoStock (SALIDA origen, ENTRADA destino) por cada producto.
    Devuelve el detalle completo para que el frontend genere el PDF.
    """
    if not transferencia.productos:
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail="Debe incluir al menos un producto en la transferencia"
        )

    if transferencia.local_origen_id == transferencia.local_destino_id:
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail="El local de origen y destino no pueden ser el mismo"
        )

    # Verificar que ambos locales existen
    local_origen = session.get(Local, transferencia.local_origen_id)
    if not local_origen:
        raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail=f"Local origen {transferencia.local_origen_id} no encontrado")

    local_destino = session.get(Local, transferencia.local_destino_id)
    if not local_destino:
        raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail=f"Local destino {transferencia.local_destino_id} no encontrado")

    # Verificar duplicados en el payload
    ids = [item.accesorio_id for item in transferencia.productos]
    if len(ids) != len(set(ids)):
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail="No se permiten accesorios duplicados en la transferencia."
        )

    try:
        # Crear registro de transferencia
        registro_transf = Transferencia(
            local_origen_id=transferencia.local_origen_id,
            local_destino_id=transferencia.local_destino_id,
            usuario_id=current_user.usuario_id,
            observaciones=transferencia.observaciones or None,
        )
        session.add(registro_transf)
        session.flush()  # obtenemos transferencia_id

        resultado_items = []

        #ItemTransferencia
        for item in transferencia.productos:
            if item.cantidad_a_transferir <= 0:
                raise HTTPException(
                    status_code=status.HTTP_400_BAD_REQUEST,
                    detail=f"La cantidad del accesorio {item.accesorio_id} debe ser mayor a cero"
                )

            accesorio = session.get(Accesorio, item.accesorio_id)
            if not accesorio:
                raise HTTPException(
                    status_code=status.HTTP_404_NOT_FOUND,
                    detail=f"Accesorio ID {item.accesorio_id} no encontrado"
                )

            # Stock origen con lock
            stock_origen_query = select(StockAccesorio).where(
                StockAccesorio.accesorio_id == item.accesorio_id,
                StockAccesorio.local_id == transferencia.local_origen_id
            ).with_for_update()
            stock_origen = session.exec(stock_origen_query).first()

            # if (transferencia.observaciones is not None):
            #     time.sleep(20) # origen

            if not stock_origen:
                raise HTTPException(
                    status_code=status.HTTP_404_NOT_FOUND,
                    detail=f"No hay stock de '{accesorio.nombre}' en el local de origen"
                )

            if stock_origen.cantidad < item.cantidad_a_transferir:
                raise HTTPException(
                    status_code=status.HTTP_400_BAD_REQUEST,
                    detail=f"Stock insuficiente de '{accesorio.nombre}'. Disponible: {stock_origen.cantidad}, solicitado: {item.cantidad_a_transferir}"
                )

            # Stock destino con lock
            stock_destino_query = select(StockAccesorio).where(
                StockAccesorio.accesorio_id == item.accesorio_id,
                StockAccesorio.local_id == transferencia.local_destino_id
            ).with_for_update()
            stock_destino = session.exec(stock_destino_query).first()


            if (transferencia.observaciones is not None):
                time.sleep(20) # destino

            # si no hay stock en el local de destino lo crea (no deberia pasar)
            if not stock_destino:
                stock_destino = StockAccesorio(
                    accesorio_id=item.accesorio_id,
                    local_id=transferencia.local_destino_id,
                    cantidad=0
                )
                session.add(stock_destino)
                session.flush()

            # if (transferencia.local_origen_id == 1):
            #     time.sleep(20) # rocca 199

            # SALIDA en origen + ENTRADA en destino
            # (cada helper aplica el delta sobre su fila y graba el movimiento con snapshots)
            mov_salida = aplicar_movimiento_stock(
                session, stock_origen,
                tipo_movimiento=TipoMovimiento.SALIDA,
                cantidad=-abs(item.cantidad_a_transferir),
                motivo=f"Transferencia #{registro_transf.transferencia_id} a {local_destino.nombre}",
                usuario_id=current_user.usuario_id,
                transferencia_id=registro_transf.transferencia_id,
            )
            mov_entrada = aplicar_movimiento_stock(
                session, stock_destino,
                tipo_movimiento=TipoMovimiento.ENTRADA,
                cantidad=item.cantidad_a_transferir,
                motivo=f"Transferencia #{registro_transf.transferencia_id} desde {local_origen.nombre}",
                usuario_id=current_user.usuario_id,
                transferencia_id=registro_transf.transferencia_id,
            )

            resultado_items.append({
                "accesorio_id": item.accesorio_id,
                "accesorio_nombre": accesorio.nombre,
                "cantidad_transferida": item.cantidad_a_transferir,
                "stock_origen_anterior": mov_salida.stock_anterior,
                "stock_origen_nuevo": mov_salida.stock_nuevo,
                "stock_destino_anterior": mov_entrada.stock_anterior,
                "stock_destino_nuevo": mov_entrada.stock_nuevo,
            })

        session.commit()
        session.refresh(registro_transf)

        return {
            "mensaje": "Transferencia de lote realizada exitosamente",
            "transferencia_id": registro_transf.transferencia_id,
            "fecha": registro_transf.fecha,
            "local_origen_id": transferencia.local_origen_id,
            "local_origen_nombre": local_origen.nombre,
            "local_destino_id": transferencia.local_destino_id,
            "local_destino_nombre": local_destino.nombre,
            "usuario_id": current_user.usuario_id,
            "observaciones": registro_transf.observaciones,
            "items": resultado_items,
        }

    except exc.IntegrityError as e:
        session.rollback()
        raise HTTPException(status_code=status.HTTP_400_BAD_REQUEST, detail=f"Error de integridad: {e.orig}")
    except HTTPException:
        session.rollback()
        raise
    except Exception as e:
        session.rollback()
        raise HTTPException(status_code=status.HTTP_500_INTERNAL_SERVER_ERROR, detail=f"Error inesperado: {str(e)}")


# ─── TRANSFERENCIA INDIVIDUAL (deprecated — mantenida por compatibilidad) ─────
@router.post("/transferir", status_code=status.HTTP_200_OK)
def transferir_stock(
    transferencia: TransferenciaStockCreate,
    session: Session = Depends(get_session),
    current_user: UsuarioActual = Depends(get_current_user)
):
    """
    Transferir un único accesorio entre dos locales.
    Deprecated: usar /transferir-lote con un solo ítem.
    """
    if transferencia.local_origen_id == transferencia.local_destino_id:
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail="El local de origen y destino no pueden ser el mismo"
        )

    try:
        stock_origen_query = select(StockAccesorio).where(
            StockAccesorio.accesorio_id == transferencia.accesorio_id,
            StockAccesorio.local_id == transferencia.local_origen_id
        ).with_for_update()
        stock_origen = session.exec(stock_origen_query).first()

        if not stock_origen:
            raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail="Stock de origen no encontrado")

        if stock_origen.cantidad < transferencia.cantidad:
            raise HTTPException(
                status_code=status.HTTP_400_BAD_REQUEST,
                detail=f"Stock insuficiente en origen. Disponible: {stock_origen.cantidad}"
            )

        stock_destino_query = select(StockAccesorio).where(
            StockAccesorio.accesorio_id == transferencia.accesorio_id,
            StockAccesorio.local_id == transferencia.local_destino_id
        ).with_for_update()
        stock_destino = session.exec(stock_destino_query).first()

        if not stock_destino:
            stock_destino = StockAccesorio(
                accesorio_id=transferencia.accesorio_id,
                local_id=transferencia.local_destino_id,
                cantidad=0
            )
            session.add(stock_destino)

        stock_origen.cantidad -= transferencia.cantidad
        stock_destino.cantidad += transferencia.cantidad

        motivo_salida = f"Transferencia a local {transferencia.local_destino_id}"
        if transferencia.motivo:
            motivo_salida += f" - {transferencia.motivo}"

        motivo_entrada = f"Transferencia desde local {transferencia.local_origen_id}"
        if transferencia.motivo:
            motivo_entrada += f" - {transferencia.motivo}"

        mov_salida = MovimientoStock(
            accesorio_id=transferencia.accesorio_id,
            local_id=transferencia.local_origen_id,
            tipo_movimiento=TipoMovimiento.SALIDA,
            cantidad=-abs(transferencia.cantidad),
            motivo=motivo_salida,
            usuario_id=current_user.usuario_id
        ) # type: ignore
        mov_entrada = MovimientoStock(
            accesorio_id=transferencia.accesorio_id,
            local_id=transferencia.local_destino_id,
            tipo_movimiento=TipoMovimiento.ENTRADA,
            cantidad=transferencia.cantidad,
            motivo=motivo_entrada,
            usuario_id=current_user.usuario_id
        ) # type: ignore

        session.add(mov_salida)
        session.add(mov_entrada)
        session.commit()
        session.refresh(mov_salida)
        session.refresh(mov_entrada)

        return {
            "mensaje": "Transferencia realizada exitosamente",
            "local_origen": {"local_id": transferencia.local_origen_id, "stock_actual": stock_origen.cantidad, "movimiento_id": mov_salida.id},
            "local_destino": {"local_id": transferencia.local_destino_id, "stock_actual": stock_destino.cantidad, "movimiento_id": mov_entrada.id},
            "cantidad_transferida": transferencia.cantidad
        }

    except exc.IntegrityError:
        session.rollback()
        raise HTTPException(status_code=status.HTTP_400_BAD_REQUEST, detail="Error en la integridad de datos")
    except HTTPException:
        session.rollback()
        raise
    except Exception as e:
        session.rollback()
        raise HTTPException(status_code=status.HTTP_500_INTERNAL_SERVER_ERROR, detail=f"Error inesperado: {str(e)}")


# ─── INGRESO-EGRESO INDIVIDUAL (deprecated — solo para compatibilidad) ─────────
@router.post("/ingreso-egreso", status_code=status.HTTP_200_OK)
def registro_movimiento_stock(
    movimiento: MovimientoStockCreate,
    session: Session = Depends(get_session),
    current_user: UsuarioActual = Depends(require_admin)
):
    """
    Movimiento individual de entrada o salida.
    Para ingresos prefer usar /ingresar-lote. Para egresos usar /egreso.
    """
    if movimiento.tipo_movimiento not in [TipoMovimiento.ENTRADA, TipoMovimiento.SALIDA]:
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail="El tipo de movimiento debe ser ENTRADA o SALIDA"
        )
    if movimiento.cantidad <= 0:
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail="La cantidad debe ser mayor a cero"
        )

    try:
        stock_query = select(StockAccesorio).where(
            StockAccesorio.accesorio_id == movimiento.accesorio_id,
            StockAccesorio.local_id == movimiento.local_id
        ).with_for_update()
        stock = session.exec(stock_query).first()

        if not stock:
            if movimiento.tipo_movimiento == TipoMovimiento.ENTRADA:
                stock = StockAccesorio(
                    accesorio_id=movimiento.accesorio_id,
                    local_id=movimiento.local_id,
                    cantidad=0
                )
                session.add(stock)
            else:
                raise HTTPException(
                    status_code=status.HTTP_400_BAD_REQUEST,
                    detail=f"No existe stock del accesorio con ID {movimiento.accesorio_id}"
                )

        if movimiento.tipo_movimiento == TipoMovimiento.ENTRADA:
            stock.cantidad += movimiento.cantidad
            motivo = f"Ingreso en local {movimiento.local_id}"
            cantidad_movida = movimiento.cantidad
        else:
            if stock.cantidad < movimiento.cantidad:
                raise HTTPException(
                    status_code=status.HTTP_400_BAD_REQUEST,
                    detail=f"Stock insuficiente. Disponible: {stock.cantidad}, Solicitado: {movimiento.cantidad}"
                )
            stock.cantidad -= movimiento.cantidad
            cantidad_movida = -abs(movimiento.cantidad)
            motivo = f"Salida en local {movimiento.local_id}"

        if movimiento.motivo:
            motivo += f" - {movimiento.motivo}"

        mov_nuevo = MovimientoStock(
            accesorio_id=movimiento.accesorio_id,
            local_id=movimiento.local_id,
            tipo_movimiento=movimiento.tipo_movimiento,
            cantidad=cantidad_movida,
            motivo=motivo,
            usuario_id=current_user.usuario_id
        ) #type: ignore

        session.add(mov_nuevo)
        session.commit()
        session.refresh(mov_nuevo)

        return {
            "mensaje": "Movimiento realizado exitosamente",
            "movimiento": {"id": mov_nuevo.id, "tipo": movimiento.tipo_movimiento.value, "cantidad": movimiento.cantidad, "fecha": mov_nuevo.fecha},
            "stock": {"local_id": movimiento.local_id, "accesorio_id": movimiento.accesorio_id, "cantidad_actual": stock.cantidad}
        }

    except exc.IntegrityError:
        session.rollback()
        raise HTTPException(status_code=status.HTTP_400_BAD_REQUEST, detail="Error en la integridad de datos")
    except HTTPException:
        session.rollback()
        raise
    except Exception as e:
        session.rollback()
        raise HTTPException(status_code=status.HTTP_500_INTERNAL_SERVER_ERROR, detail=f"Error inesperado: {str(e)}")


@router.post("/seed-stock")
def seed_stock(
    cantidad: int = 100,
    sumar: bool = True,
    session: Session = Depends(get_session),
    current_user: UsuarioActual = Depends(require_admin)
):
    """
    Ajusta el stock de todos los accesorios activos en todos los locales.
    - sumar=true (default): suma `cantidad` al stock actual.
    - sumar=false: setea el stock exactamente a `cantidad`.
    Registra un movimiento AJUSTE por cada fila modificada.
    """
    stocks = session.exec(
        select(StockAccesorio)
        .join(Accesorio, StockAccesorio.accesorio_id == Accesorio.accesorio_id)  # type: ignore
        .where(Accesorio.activo == True)  # type: ignore
    ).all()

    for stock in stocks:
        cantidad_anterior = stock.cantidad
        cantidad_nueva = cantidad_anterior + cantidad if sumar else cantidad
        diferencia = cantidad_nueva - cantidad_anterior

        if diferencia != 0:
            aplicar_movimiento_stock(
                session, stock,
                tipo_movimiento=TipoMovimiento.AJUSTE,
                cantidad=diferencia,
                motivo=f"Seed stock ({'suma' if sumar else 'seteo'} {cantidad})",
                usuario_id=current_user.usuario_id,
            )

    session.commit()
    return {
        "mensaje": f"Stock {'sumado' if sumar else 'seteado'} a {cantidad} en {len(stocks)} registros.",
        "registros_actualizados": len(stocks),
    }
