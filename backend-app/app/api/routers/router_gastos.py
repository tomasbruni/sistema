from datetime import date, datetime
from decimal import Decimal, ROUND_HALF_UP
from io import BytesIO
from typing import Optional

import calendar
import re

import openpyxl
from fastapi import APIRouter, Depends, HTTPException, Query, status
from fastapi.responses import StreamingResponse
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from sqlalchemy import exc
from sqlmodel import Session, SQLModel, select, func

from app.db.session import get_session
from app.db.models import Gasto, TipoGasto, TipoFactura, Usuario, ProveedorFiscal
from app.api.deps import get_current_user, require_admin, UsuarioActual
from app.api.funciones.fechas import start_of_day, end_of_day, TZ_AR
from app.api.funciones.gastos_funciones import primer_dia_mes, condicion_rango_gastos


router = APIRouter(prefix="/gastos", tags=["Gastos reales y facturados"])

CERO = Decimal("0.00")
DOS_DECIMALES = Decimal("0.01")

# Columnas del libro IVA compras, en el orden del Excel del contador (G..R).
# El total del comprobante es SIEMPRE la suma de todas.
NETOS = ("neto_105", "neto_21", "neto_27")
IVAS = ("iva_105", "iva_21", "iva_27")
COLUMNAS_IMPORTE = (
    *NETOS, *IVAS,
    "exento", "percepcion_iva", "percepcion_iibb_bsas", "percepcion_iibb_caba",
    "otros_impuestos", "no_gravado",
)
# En factura C no hay IVA discriminado: estas columnas tienen que ir en 0.
COLUMNAS_SOLO_A = (*NETOS, *IVAS, "percepcion_iva")
# Columna donde va el importe de una factura C.
# PENDIENTE confirmar con el contador (el Excel de referencia no tiene facturas C).
COLUMNA_IMPORTE_C = "no_gravado"


# ── Schemas ──────────────────────────────────────────────────────────────────

class GastoBase(SQLModel):
    descripcion: Optional[str] = None
    total: Optional[Decimal] = None   # REAL: plata que salio | factura: total impreso del comprobante
    fecha: Optional[date] = None      # factura: fecha de emision
    periodo: Optional[date] = None    # factura: mes de imputacion (se normaliza al dia 1)
    tipo_factura: Optional[TipoFactura] = None
    proveedor_fiscal_id: Optional[int] = None
    punto_venta: Optional[str] = None
    numero_comprobante: Optional[str] = None
    comprada: Optional[bool] = None
    porcentaje_real: Optional[int] = None
    neto_105: Optional[Decimal] = None
    neto_21: Optional[Decimal] = None
    neto_27: Optional[Decimal] = None
    iva_105: Optional[Decimal] = None
    iva_21: Optional[Decimal] = None
    iva_27: Optional[Decimal] = None
    exento: Optional[Decimal] = None
    percepcion_iva: Optional[Decimal] = None
    percepcion_iibb_bsas: Optional[Decimal] = None
    percepcion_iibb_caba: Optional[Decimal] = None
    otros_impuestos: Optional[Decimal] = None
    no_gravado: Optional[Decimal] = None


class GastoCreate(GastoBase):
    tipo: TipoGasto
    usuario_id: Optional[int] = None


class GastoUpdate(GastoBase):
    tipo: Optional[TipoGasto] = None


class GastoRead(SQLModel):
    id: int
    tipo: TipoGasto
    descripcion: Optional[str]
    total: Decimal
    fecha: Optional[datetime]
    periodo: Optional[date]
    usuario_id: int
    tipo_factura: Optional[TipoFactura]
    proveedor_fiscal_id: Optional[int]
    proveedor_razon_social: Optional[str] = None
    proveedor_cuit: Optional[str] = None
    punto_venta: Optional[str]
    numero_comprobante: Optional[str]
    comprada: bool
    porcentaje_real: Optional[int]
    neto_105: Decimal
    neto_21: Decimal
    neto_27: Decimal
    iva_105: Decimal
    iva_21: Decimal
    iva_27: Decimal
    exento: Decimal
    percepcion_iva: Decimal
    percepcion_iibb_bsas: Decimal
    percepcion_iibb_caba: Decimal
    otros_impuestos: Decimal
    no_gravado: Decimal
    aporte_real: Decimal
    aporte_blanco: Decimal
    aporte_iva: Decimal
    aporte_percepcion_iibb: Decimal


class GastoTotales(SQLModel):
    total_real: Decimal
    total_blanco: Decimal
    iva_a_favor: Decimal
    percepciones_iibb: Decimal
    cantidad: int


# ── Logica de negocio ─────────────────────────────────────────────────────────

def _dec(v) -> Decimal:
    return CERO if v is None else Decimal(v).quantize(DOS_DECIMALES, ROUND_HALF_UP)


def _fecha_ar(dt: Optional[datetime]) -> Optional[date]:
    return dt.astimezone(TZ_AR).date() if dt else None


def _texto(v) -> Optional[str]:
    if v is None:
        return None
    v = str(v).strip()
    return v or None


def _normalizar_comprobante(v) -> Optional[str]:
    """Punto de venta / numero de comprobante en una forma unica, para que el control
    de duplicados y el Excel vean el mismo valor: sin ceros a la izquierda en cada
    tramo numerico separado por guion ("00005" -> "5", "07008-00041244" -> "7008-41244").
    Si algun tramo no es puramente numerico, queda como vino (solo sin espacios)."""
    texto = _texto(v)
    if texto is None:
        return None
    tramos = [t.strip() for t in texto.split("-")]
    if all(re.fullmatch(r"[0-9]+", t) for t in tramos):
        return "-".join(str(int(t)) for t in tramos)
    return texto


def _normalizar_y_calcular(data: dict, session: Session) -> dict:
    """Valida un estado logico completo de gasto y devuelve los campos a persistir
    (importes normalizados + aportes calculados). Sirve tanto para crear como para
    actualizar (mergeando el payload sobre el registro existente)."""
    tipo: TipoGasto = data["tipo"]
    descripcion = _texto(data.get("descripcion"))

    campos: dict = {
        "tipo": tipo,
        "descripcion": descripcion,
        "periodo": None,
        "tipo_factura": None,
        "proveedor_fiscal_id": None,
        "punto_venta": None,
        "numero_comprobante": None,
        "comprada": False,
        "porcentaje_real": None,
        **{c: CERO for c in COLUMNAS_IMPORTE},
    }

    if data.get("total") is None:
        raise HTTPException(422, "Falta el 'total'" if tipo == TipoGasto.REAL
                            else "Falta el 'total' del comprobante (tal como figura impreso)")
    total = _dec(data["total"])
    if total <= 0:
        raise HTTPException(422, "El total tiene que ser mayor a 0")
    campos["total"] = total

    if tipo == TipoGasto.REAL:
        if not descripcion:
            raise HTTPException(422, "'descripcion' es obligatoria")
        campos.update(aporte_real=total, aporte_blanco=CERO, aporte_iva=CERO,
                      aporte_percepcion_iibb=CERO)
        return campos

    # ── FACTURA ──
    tipo_factura = data.get("tipo_factura")
    if tipo_factura not in (TipoFactura.A, TipoFactura.C):
        raise HTTPException(422, "Una factura necesita 'tipo_factura' (A o C)")

    proveedor_id = data.get("proveedor_fiscal_id")
    if proveedor_id is None:
        raise HTTPException(422, "Una factura necesita 'proveedor_fiscal_id'")
    proveedor = session.get(ProveedorFiscal, proveedor_id)
    if not proveedor or not proveedor.activo:
        raise HTTPException(404, "Proveedor no encontrado o inactivo")

    punto_venta = _normalizar_comprobante(data.get("punto_venta"))
    numero = _normalizar_comprobante(data.get("numero_comprobante"))
    if not punto_venta or not numero:
        raise HTTPException(422, "Una factura necesita 'punto_venta' y 'numero_comprobante'")
    fecha = data.get("fecha")
    if fecha is None:
        raise HTTPException(422, "Una factura necesita 'fecha' (de emisión)")
    # El credito fiscal se puede imputar a un periodo posterior al de emision
    # (facturas que llegan tarde), nunca a uno anterior.
    periodo = primer_dia_mes(data.get("periodo") or fecha)
    if periodo < primer_dia_mes(fecha):
        raise HTTPException(422, "El período de imputación no puede ser anterior al mes de emisión")

    importes = {c: _dec(data.get(c)) for c in COLUMNAS_IMPORTE}
    negativos = [c for c, v in importes.items() if v < 0]
    if negativos:
        raise HTTPException(422, f"Importes negativos no permitidos: {', '.join(negativos)}")

    if tipo_factura == TipoFactura.A:
        if not any(importes[c] > 0 for c in NETOS):
            raise HTTPException(422, "La factura A necesita al menos un neto gravado")
    else:
        con_valor = [c for c in COLUMNAS_SOLO_A if importes[c] != 0]
        if con_valor:
            raise HTTPException(422, f"La factura C no discrimina IVA: {', '.join(con_valor)} tiene(n) que ir en 0")

    # El total impreso tiene que cerrar con la suma de columnas: evita cargar un
    # mismo importe en dos columnas (ej. otros impuestos y no gravado).
    suma = sum(importes.values(), CERO)
    if suma != total:
        raise HTTPException(
            422,
            f"El total declarado ({total}) no coincide con la suma de columnas ({suma}): "
            f"diferencia {total - suma}",
        )

    campos.update(
        periodo=periodo,
        tipo_factura=tipo_factura,
        proveedor_fiscal_id=proveedor_id,
        punto_venta=punto_venta,
        numero_comprobante=numero,
        **importes,
    )

    comprada = bool(data.get("comprada")) and tipo_factura == TipoFactura.A
    campos["comprada"] = comprada
    if comprada:
        pct = data.get("porcentaje_real")
        if pct is None:
            raise HTTPException(422, "La factura A comprada necesita 'porcentaje_real'")
        if pct < 0 or pct > 100:
            raise HTTPException(422, "'porcentaje_real' debe estar entre 0 y 100")
        campos["porcentaje_real"] = pct
        aporte_real = (total * Decimal(pct) / Decimal(100)).quantize(DOS_DECIMALES, ROUND_HALF_UP)
    else:
        aporte_real = CERO

    i = importes
    campos["aporte_real"] = aporte_real
    campos["aporte_blanco"] = (sum((i[c] for c in NETOS), CERO)
                               + i["exento"] + i["otros_impuestos"] + i["no_gravado"])
    campos["aporte_iva"] = sum((i[c] for c in IVAS), CERO) + i["percepcion_iva"]
    campos["aporte_percepcion_iibb"] = i["percepcion_iibb_bsas"] + i["percepcion_iibb_caba"]
    return campos


def _a_read(gasto: Gasto, proveedor: Optional[ProveedorFiscal]) -> GastoRead:
    return GastoRead(
        **gasto.model_dump(),
        proveedor_razon_social=proveedor.razon_social if proveedor else None,
        proveedor_cuit=proveedor.cuit if proveedor else None,
    )


def _guardar(session: Session, gasto: Gasto) -> GastoRead:
    session.add(gasto)
    try:
        session.commit()
    except exc.IntegrityError:
        session.rollback()
        raise HTTPException(400, "Ya existe una factura de ese proveedor con el mismo tipo, punto de venta y número")
    session.refresh(gasto)
    proveedor = session.get(ProveedorFiscal, gasto.proveedor_fiscal_id) if gasto.proveedor_fiscal_id else None
    return _a_read(gasto, proveedor)


# ── Endpoints ─────────────────────────────────────────────────────────────────

@router.post("/", response_model=GastoRead, status_code=status.HTTP_201_CREATED)
def crear_gasto(
    payload: GastoCreate,
    session: Session = Depends(get_session),
    current_user: UsuarioActual = Depends(require_admin),
):
    campos = _normalizar_y_calcular(payload.model_dump(), session)

    usuario_id = payload.usuario_id or current_user.usuario_id
    if payload.usuario_id is not None and not session.get(Usuario, payload.usuario_id):
        raise HTTPException(404, "Usuario no existe")

    gasto = Gasto(**campos, usuario_id=usuario_id)
    if payload.fecha is not None:
        gasto.fecha = start_of_day(payload.fecha)
    return _guardar(session, gasto)


@router.get("/", response_model=list[GastoRead])
def listar_gastos(
    fecha: Optional[date] = Query(default=None, description="Dia exacto (YYYY-MM-DD)"),
    fecha_desde: Optional[date] = Query(default=None),
    fecha_hasta: Optional[date] = Query(default=None),
    tipo: Optional[TipoGasto] = Query(default=None),
    tipo_factura: Optional[TipoFactura] = Query(default=None),
    proveedor_fiscal_id: Optional[int] = Query(default=None),
    offset: int = Query(default=0, ge=0),
    limit: int = Query(default=100, ge=1, le=500),
    session: Session = Depends(get_session),
    current_user: UsuarioActual = Depends(require_admin),
):
    query = (
        select(Gasto, ProveedorFiscal)
        .join(ProveedorFiscal, Gasto.proveedor_fiscal_id == ProveedorFiscal.proveedor_fiscal_id,  # type: ignore
              isouter=True)
    )
    query = _aplicar_filtros(query, fecha, fecha_desde, fecha_hasta, tipo, tipo_factura, proveedor_fiscal_id)
    query = query.order_by(Gasto.fecha.desc(), Gasto.id.desc()).offset(offset).limit(limit)  # type: ignore
    return [_a_read(g, p) for g, p in session.exec(query).all()]


@router.get("/totales", response_model=GastoTotales)
def totales_gastos(
    fecha: Optional[date] = Query(default=None, description="Dia exacto (YYYY-MM-DD)"),
    fecha_desde: Optional[date] = Query(default=None),
    fecha_hasta: Optional[date] = Query(default=None),
    tipo: Optional[TipoGasto] = Query(default=None),
    tipo_factura: Optional[TipoFactura] = Query(default=None),
    proveedor_fiscal_id: Optional[int] = Query(default=None),
    session: Session = Depends(get_session),
    current_user: UsuarioActual = Depends(require_admin),
):
    query = select(
        func.coalesce(func.sum(Gasto.aporte_real), 0),
        func.coalesce(func.sum(Gasto.aporte_blanco), 0),
        func.coalesce(func.sum(Gasto.aporte_iva), 0),
        func.coalesce(func.sum(Gasto.aporte_percepcion_iibb), 0),
        func.count(Gasto.id), #type: ignore
    )
    query = _aplicar_filtros(query, fecha, fecha_desde, fecha_hasta, tipo, tipo_factura, proveedor_fiscal_id)
    total_real, total_blanco, iva_a_favor, percep_iibb, cantidad = session.exec(query).one()
    return GastoTotales(
        total_real=Decimal(total_real),
        total_blanco=Decimal(total_blanco),
        iva_a_favor=Decimal(iva_a_favor),
        percepciones_iibb=Decimal(percep_iibb),
        cantidad=cantidad,
    )


@router.get("/iva-compras/excel")
def iva_compras_excel(
    periodo: str = Query(description="Mes de imputación (YYYY-MM)"),
    session: Session = Depends(get_session),
    current_user: UsuarioActual = Depends(require_admin),
):
    """Libro IVA compras de un periodo, con el formato de la hoja COMPRAS del contador.
    Solo facturas (los gastos REAL no tienen comprobante), filtradas por periodo de
    imputacion; la columna FECHA muestra la fecha de emision."""
    m = re.fullmatch(r"(\d{4})-(\d{2})", periodo)
    if not m or not 1 <= int(m.group(2)) <= 12:
        raise HTTPException(422, "'periodo' debe tener formato YYYY-MM")
    anio, mes = int(m.group(1)), int(m.group(2))
    inicio = date(anio, mes, 1)
    fin = date(anio, mes, calendar.monthrange(anio, mes)[1])

    filas = session.exec(
        select(Gasto, ProveedorFiscal)
        .join(ProveedorFiscal, Gasto.proveedor_fiscal_id == ProveedorFiscal.proveedor_fiscal_id)  # type: ignore
        .where(Gasto.tipo == TipoGasto.FACTURA)
        .where(Gasto.periodo == inicio)
        .order_by(Gasto.fecha, Gasto.id)  # type: ignore
    ).all()

    xlsx_bytes = _build_iva_compras_excel(list(filas), fin)
    return StreamingResponse(
        BytesIO(xlsx_bytes),
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="iva_compras_{periodo}.xlsx"'},
    )


@router.get("/{gasto_id}", response_model=GastoRead)
def obtener_gasto(
    gasto_id: int,
    session: Session = Depends(get_session),
    current_user: UsuarioActual = Depends(require_admin),
):
    gasto = session.get(Gasto, gasto_id)
    if not gasto:
        raise HTTPException(status_code=404, detail="Gasto no encontrado")
    proveedor = session.get(ProveedorFiscal, gasto.proveedor_fiscal_id) if gasto.proveedor_fiscal_id else None
    return _a_read(gasto, proveedor)


@router.patch("/{gasto_id}", response_model=GastoRead)
def actualizar_gasto(
    gasto_id: int,
    payload: GastoUpdate,
    session: Session = Depends(get_session),
    current_user: UsuarioActual = Depends(require_admin),
):
    gasto = session.get(Gasto, gasto_id)
    if not gasto:
        raise HTTPException(status_code=404, detail="Gasto no encontrado")

    cambios = payload.model_dump(exclude_unset=True)
    fecha = cambios.pop("fecha", None)

    # Merge del estado actual con los cambios y re-validacion / recalculo completo.
    fecha_anterior = _fecha_ar(gasto.fecha)
    estado = gasto.model_dump()
    estado["fecha"] = fecha or fecha_anterior
    # Si cambia la emision y el periodo nunca se toco (era el mes de emision),
    # el periodo acompaña a la fecha nueva.
    if (fecha is not None and "periodo" not in cambios and fecha_anterior is not None
            and gasto.periodo == primer_dia_mes(fecha_anterior)):
        estado["periodo"] = None
    estado.update(cambios)

    campos = _normalizar_y_calcular(estado, session)
    for campo, valor in campos.items():
        setattr(gasto, campo, valor)
    if fecha is not None:
        gasto.fecha = start_of_day(fecha)
    return _guardar(session, gasto)


@router.delete("/{gasto_id}", status_code=status.HTTP_204_NO_CONTENT)
def eliminar_gasto(
    gasto_id: int,
    session: Session = Depends(get_session),
    current_user: UsuarioActual = Depends(require_admin),
):
    gasto = session.get(Gasto, gasto_id)
    if not gasto:
        raise HTTPException(status_code=404, detail="Gasto no encontrado")
    session.delete(gasto)
    session.commit()


# ── Helpers ───────────────────────────────────────────────────────────────────

def _aplicar_filtros(query, fecha, fecha_desde, fecha_hasta, tipo, tipo_factura, proveedor_fiscal_id=None):
    if fecha:
        query = query.where(
            Gasto.fecha >= start_of_day(fecha),  # type: ignore
            Gasto.fecha <= end_of_day(fecha),    # type: ignore
        )
    elif fecha_desde or fecha_hasta:
        # Facturas por periodo de imputacion, gastos REAL por fecha.
        query = query.where(condicion_rango_gastos(fecha_desde, fecha_hasta))
    if tipo is not None:
        query = query.where(Gasto.tipo == tipo)
    if tipo_factura is not None:
        query = query.where(Gasto.tipo_factura == tipo_factura)
    if proveedor_fiscal_id is not None:
        query = query.where(Gasto.proveedor_fiscal_id == proveedor_fiscal_id)
    return query


# ── Excel IVA compras ─────────────────────────────────────────────────────────
# Replica la hoja COMPRAS de "IVA COMPRAS VENTAS.xlsx" del contador.

_IVA_ENCABEZADOS = [
    "FECHA", "TIPO      FACTURA", "NOMBRE", "CUIT", "SUCURSAL", "N° COMPROBANTE",
    "NETO 10,5%", "NETO 21%", "NETO 27%", "IVA 10,5%", "IVA 21%", "IVA 27%",
    "EXENTO", "PERCEP IVA", "PERCEP IIBB", "PERCEP IIBB", "OTROS IMP.", "NO GRAV.",
    "TOTAL", "CONCEPTOS",
]
_IVA_ANCHOS = [11, 9, 38, 15, 10, 15] + [13] * 12 + [14, 30]
_IVA_PRIMERA_COL_IMPORTE = 7   # G
_IVA_COL_TOTAL = 19            # S

_ARIAL = "Arial"
_THIN = Side(style="thin")
_BORDE = Border(left=_THIN, right=_THIN, top=_THIN, bottom=_THIN)
_FILL_TITULO = PatternFill("solid", fgColor="FEE1CC")
_FILL_GRIS = PatternFill("solid", fgColor="D9D9D9")
_CENTRO = Alignment(horizontal="center", vertical="center", wrap_text=True)
_FMT_IMPORTE = "#,##0.00"


def _num_o_str(v: Optional[str]):
    """Sucursal / comprobante: numero si es puramente numerico (como en el original)."""
    if v is None:
        return None
    return int(v) if v.isdigit() else v


def _build_iva_compras_excel(filas: list[tuple[Gasto, ProveedorFiscal]], hasta: date) -> bytes:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "COMPRAS"  # type: ignore

    ultima_col = len(_IVA_ENCABEZADOS)

    # Titulo (A1:T2) y fila 3 (fecha del periodo + jurisdicciones de IIBB)
    ws.merge_cells(start_row=1, start_column=1, end_row=2, end_column=ultima_col)
    c = ws.cell(row=1, column=1, value="IVA COMPRAS ")
    c.font = Font(name=_ARIAL, size=18, bold=True)
    c.fill = _FILL_TITULO
    c.alignment = _CENTRO

    ws.merge_cells(start_row=3, start_column=5, end_row=3, end_column=6)
    c = ws.cell(row=3, column=5, value=hasta)
    c.number_format = "dd-mm-yy"
    for col, texto in ((5, None), (15, "BS AS"), (16, "CAPITAL")):
        c = ws.cell(row=3, column=col)
        if texto:
            c.value = texto
        c.font = Font(name=_ARIAL, size=10)
        c.fill = _FILL_GRIS
        c.alignment = _CENTRO
        c.border = _BORDE

    for col, texto in enumerate(_IVA_ENCABEZADOS, start=1):
        c = ws.cell(row=4, column=col, value=texto)
        c.font = Font(name=_ARIAL, size=10, bold=True)
        c.fill = _FILL_GRIS
        c.alignment = _CENTRO
        c.border = _BORDE
    ws.row_dimensions[4].height = 27

    fila = 5
    for g, p in filas:
        valores = [
            _fecha_ar(g.fecha),
            g.tipo_factura.value if g.tipo_factura else None,
            p.razon_social,
            p.cuit,
            _num_o_str(g.punto_venta),
            _num_o_str(g.numero_comprobante),
            *[float(getattr(g, col)) or None for col in COLUMNAS_IMPORTE],
            f"=SUM({get_column_letter(_IVA_PRIMERA_COL_IMPORTE)}{fila}:"
            f"{get_column_letter(_IVA_COL_TOTAL - 1)}{fila})",
            g.descripcion,
        ]
        for col, v in enumerate(valores, start=1):
            c = ws.cell(row=fila, column=col, value=v)
            c.font = Font(name=_ARIAL, size=10)
            c.border = _BORDE
            if col == 1:
                c.number_format = "d/mm/yyyy"
                c.alignment = Alignment(horizontal="center")
            elif col in (2, 5, 6):
                c.alignment = Alignment(horizontal="center")
            elif _IVA_PRIMERA_COL_IMPORTE <= col <= _IVA_COL_TOTAL:
                c.number_format = _FMT_IMPORTE
        fila += 1

    # Fila de totales
    c = ws.cell(row=fila, column=1, value="TOTALES")
    ws.merge_cells(start_row=fila, start_column=1, end_row=fila, end_column=_IVA_PRIMERA_COL_IMPORTE - 1)
    for col in range(1, ultima_col + 1):
        c = ws.cell(row=fila, column=col)
        c.font = Font(name=_ARIAL, size=10, bold=True)
        c.fill = _FILL_GRIS
        c.border = _BORDE
        if col == 1:
            c.alignment = _CENTRO
        if _IVA_PRIMERA_COL_IMPORTE <= col <= _IVA_COL_TOTAL:
            letra = get_column_letter(col)
            c.value = f"=SUM({letra}5:{letra}{fila - 1})" if fila > 5 else 0
            c.number_format = _FMT_IMPORTE

    for col, ancho in enumerate(_IVA_ANCHOS, start=1):
        ws.column_dimensions[get_column_letter(col)].width = ancho
    ws.freeze_panes = "A5"

    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()
