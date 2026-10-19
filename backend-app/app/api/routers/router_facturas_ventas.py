from datetime import date
from decimal import Decimal
from io import BytesIO
from typing import Optional

import openpyxl
from fastapi import APIRouter, Depends, HTTPException, Query, status
from fastapi.responses import StreamingResponse
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from sqlalchemy import exc
from sqlmodel import Session, SQLModel, select, func

from app.db.session import get_session
from app.db.models import FacturaVenta, TipoFacturaVenta, Usuario
from app.api.deps import require_admin, UsuarioActual
from app.api.funciones.gastos_funciones import (
    CERO, dec, texto_o_none, normalizar_comprobante, rango_mes,
)


router = APIRouter(prefix="/facturas-ventas", tags=["Facturas de venta (IVA ventas)"])

# Columnas del libro IVA ventas, en el orden de la hoja VENTAS del contador.
# El total del comprobante es SIEMPRE neto + IVA.
COLUMNAS_IMPORTE = ("neto_21", "iva_21")


# ── Schemas ──────────────────────────────────────────────────────────────────

class FacturaVentaBase(SQLModel):
    descripcion: Optional[str] = None
    fecha: Optional[date] = None      # fecha de emision
    punto_venta: Optional[str] = None
    numero_comprobante: Optional[str] = None
    neto_21: Optional[Decimal] = None
    iva_21: Optional[Decimal] = None
    total: Optional[Decimal] = None   # total impreso del comprobante


class FacturaVentaCreate(FacturaVentaBase):
    tipo_factura: TipoFacturaVenta
    usuario_id: Optional[int] = None


class FacturaVentaUpdate(FacturaVentaBase):
    tipo_factura: Optional[TipoFacturaVenta] = None


class FacturaVentaRead(SQLModel):
    id: int
    tipo_factura: TipoFacturaVenta
    punto_venta: str
    numero_comprobante: str
    fecha: date
    descripcion: Optional[str]
    usuario_id: int
    neto_21: Decimal
    iva_21: Decimal
    total: Decimal


class FacturaVentaTotales(SQLModel):
    neto_21: Decimal
    iva_21: Decimal
    total: Decimal
    cantidad: int


# ── Logica de negocio ─────────────────────────────────────────────────────────

def _normalizar_y_validar(data: dict) -> dict:
    """Valida un estado logico completo de factura y devuelve los campos a persistir.
    Sirve tanto para crear como para actualizar (mergeando el payload sobre el
    registro existente)."""
    tipo_factura = data.get("tipo_factura")
    if tipo_factura not in (TipoFacturaVenta.A, TipoFacturaVenta.B):
        raise HTTPException(422, "La factura necesita 'tipo_factura' (A o B)")

    punto_venta = normalizar_comprobante(data.get("punto_venta"))
    numero = normalizar_comprobante(data.get("numero_comprobante"))
    if not punto_venta or not numero:
        raise HTTPException(422, "La factura necesita 'punto_venta' y 'numero_comprobante'")
    fecha = data.get("fecha")
    if fecha is None:
        raise HTTPException(422, "La factura necesita 'fecha' (de emisión)")

    if data.get("total") is None:
        raise HTTPException(422, "Falta el 'total' del comprobante (tal como figura impreso)")
    total = dec(data["total"])
    if total <= 0:
        raise HTTPException(422, "El total tiene que ser mayor a 0")

    importes = {c: dec(data.get(c)) for c in COLUMNAS_IMPORTE}
    negativos = [c for c, v in importes.items() if v < 0]
    if negativos:
        raise HTTPException(422, f"Importes negativos no permitidos: {', '.join(negativos)}")
    if importes["neto_21"] <= 0:
        raise HTTPException(422, "La factura necesita el neto gravado 21%")

    suma = sum(importes.values(), CERO)
    if suma != total:
        raise HTTPException(
            422,
            f"El total declarado ({total}) no coincide con neto + IVA ({suma}): "
            f"diferencia {total - suma}",
        )

    return {
        "tipo_factura": tipo_factura,
        "punto_venta": punto_venta,
        "numero_comprobante": numero,
        "fecha": fecha,
        "descripcion": texto_o_none(data.get("descripcion")),
        "total": total,
        **importes,
    }


def _guardar(session: Session, factura: FacturaVenta) -> FacturaVentaRead:
    session.add(factura)
    try:
        session.commit()
    except exc.IntegrityError:
        session.rollback()
        raise HTTPException(400, "Ya existe una factura con el mismo tipo, punto de venta y número")
    session.refresh(factura)
    return FacturaVentaRead(**factura.model_dump())


# ── Endpoints ─────────────────────────────────────────────────────────────────

@router.post("/", response_model=FacturaVentaRead, status_code=status.HTTP_201_CREATED)
def crear_factura_venta(
    payload: FacturaVentaCreate,
    session: Session = Depends(get_session),
    current_user: UsuarioActual = Depends(require_admin),
):
    campos = _normalizar_y_validar(payload.model_dump())

    usuario_id = payload.usuario_id or current_user.usuario_id
    if payload.usuario_id is not None and not session.get(Usuario, payload.usuario_id):
        raise HTTPException(404, "Usuario no existe")

    return _guardar(session, FacturaVenta(**campos, usuario_id=usuario_id))


@router.get("/", response_model=list[FacturaVentaRead])
def listar_facturas_ventas(
    fecha_desde: Optional[date] = Query(default=None),
    fecha_hasta: Optional[date] = Query(default=None),
    tipo_factura: Optional[TipoFacturaVenta] = Query(default=None),
    offset: int = Query(default=0, ge=0),
    limit: int = Query(default=100, ge=1, le=500),
    session: Session = Depends(get_session),
    current_user: UsuarioActual = Depends(require_admin),
):
    query = _aplicar_filtros(select(FacturaVenta), fecha_desde, fecha_hasta, tipo_factura)
    query = query.order_by(FacturaVenta.fecha.desc(), FacturaVenta.id.desc()).offset(offset).limit(limit)  # type: ignore
    return [FacturaVentaRead(**f.model_dump()) for f in session.exec(query).all()]


@router.get("/totales", response_model=FacturaVentaTotales)
def totales_facturas_ventas(
    fecha_desde: Optional[date] = Query(default=None),
    fecha_hasta: Optional[date] = Query(default=None),
    tipo_factura: Optional[TipoFacturaVenta] = Query(default=None),
    session: Session = Depends(get_session),
    current_user: UsuarioActual = Depends(require_admin),
):
    query = select(
        func.coalesce(func.sum(FacturaVenta.neto_21), 0),
        func.coalesce(func.sum(FacturaVenta.iva_21), 0),
        func.coalesce(func.sum(FacturaVenta.total), 0),
        func.count(FacturaVenta.id),  # type: ignore
    )
    query = _aplicar_filtros(query, fecha_desde, fecha_hasta, tipo_factura)
    neto, iva, total, cantidad = session.exec(query).one()
    return FacturaVentaTotales(
        neto_21=Decimal(neto),
        iva_21=Decimal(iva),
        total=Decimal(total),
        cantidad=cantidad,
    )


@router.get("/iva-ventas/excel")
def iva_ventas_excel(
    periodo: str = Query(description="Mes de emisión (YYYY-MM)"),
    session: Session = Depends(get_session),
    current_user: UsuarioActual = Depends(require_admin),
):
    """Libro IVA ventas de un mes (por fecha de emision), con el formato de la hoja
    VENTAS del contador."""
    inicio, fin = rango_mes(periodo)
    facturas = session.exec(
        _aplicar_filtros(select(FacturaVenta), inicio, fin, None)
        .order_by(FacturaVenta.fecha, FacturaVenta.id)  # type: ignore
    ).all()

    xlsx_bytes = _build_iva_ventas_excel(list(facturas))
    return StreamingResponse(
        BytesIO(xlsx_bytes),
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="iva_ventas_{periodo}.xlsx"'},
    )


@router.get("/{factura_id}", response_model=FacturaVentaRead)
def obtener_factura_venta(
    factura_id: int,
    session: Session = Depends(get_session),
    current_user: UsuarioActual = Depends(require_admin),
):
    factura = session.get(FacturaVenta, factura_id)
    if not factura:
        raise HTTPException(status_code=404, detail="Factura no encontrada")
    return FacturaVentaRead(**factura.model_dump())


@router.patch("/{factura_id}", response_model=FacturaVentaRead)
def actualizar_factura_venta(
    factura_id: int,
    payload: FacturaVentaUpdate,
    session: Session = Depends(get_session),
    current_user: UsuarioActual = Depends(require_admin),
):
    factura = session.get(FacturaVenta, factura_id)
    if not factura:
        raise HTTPException(status_code=404, detail="Factura no encontrada")

    # Merge del estado actual con los cambios y re-validacion completa.
    estado = factura.model_dump()
    estado.update(payload.model_dump(exclude_unset=True))
    for campo, valor in _normalizar_y_validar(estado).items():
        setattr(factura, campo, valor)
    return _guardar(session, factura)


@router.delete("/{factura_id}", status_code=status.HTTP_204_NO_CONTENT)
def eliminar_factura_venta(
    factura_id: int,
    session: Session = Depends(get_session),
    current_user: UsuarioActual = Depends(require_admin),
):
    factura = session.get(FacturaVenta, factura_id)
    if not factura:
        raise HTTPException(status_code=404, detail="Factura no encontrada")
    session.delete(factura)
    session.commit()


# ── Helpers ───────────────────────────────────────────────────────────────────

def _aplicar_filtros(query, fecha_desde, fecha_hasta, tipo_factura):
    if fecha_desde:
        query = query.where(FacturaVenta.fecha >= fecha_desde)  # type: ignore
    if fecha_hasta:
        query = query.where(FacturaVenta.fecha <= fecha_hasta)  # type: ignore
    if tipo_factura is not None:
        query = query.where(FacturaVenta.tipo_factura == tipo_factura)
    return query


# ── Excel IVA ventas ──────────────────────────────────────────────────────────
# Replica la hoja VENTAS de "IVA COMPRAS VENTAS.xlsx" del contador.

_IVA_ENCABEZADOS = [
    "FECHA", "TIPO DE FACTURA", "SUCURSAL", "N° DE COMPROBANTE",
    "NETO 21%", "IVA 21%", "TOTAL", "CONCEPTOS",
]
_IVA_ANCHOS = [14.5, 10, 11, 17, 14.5, 14.5, 14.5, 30]
_IVA_PRIMERA_COL_IMPORTE = 5   # E
_IVA_COL_TOTAL = 7             # G

_ARIAL = "Arial"
_THIN = Side(style="thin")
_BORDE = Border(left=_THIN, right=_THIN, top=_THIN, bottom=_THIN)
_FILL_TITULO = PatternFill("solid", fgColor="FEE1CC")
_FILL_GRIS = PatternFill("solid", fgColor="D9D9D9")
_CENTRO = Alignment(horizontal="center", vertical="center", wrap_text=True)
_FMT_IMPORTE = "#,##0.00"


def _num_o_str(v: Optional[str]):
    """Sucursal / comprobante: numero si es puramente numerico (como en el original)."""
    if v is None:
        return None
    return int(v) if v.isdigit() else v


def _build_iva_ventas_excel(facturas: list[FacturaVenta]) -> bytes:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "VENTAS"  # type: ignore

    ultima_col = len(_IVA_ENCABEZADOS)

    # Titulo (A1:H2) y fila 3 ("FACTURA NUMERO" sobre sucursal + comprobante)
    ws.merge_cells(start_row=1, start_column=1, end_row=2, end_column=ultima_col)
    c = ws.cell(row=1, column=1, value="IVA VENTAS ")
    c.font = Font(name=_ARIAL, size=18, bold=True)
    c.fill = _FILL_TITULO
    c.alignment = _CENTRO

    ws.merge_cells(start_row=3, start_column=3, end_row=3, end_column=4)
    for col, texto in ((3, "FACTURA NUMERO"), (4, None)):
        c = ws.cell(row=3, column=col)
        if texto:
            c.value = texto
        c.font = Font(name=_ARIAL, size=10, bold=True)
        c.fill = _FILL_GRIS
        c.alignment = _CENTRO
        c.border = _BORDE

    for col, texto in enumerate(_IVA_ENCABEZADOS, start=1):
        c = ws.cell(row=4, column=col, value=texto)
        c.font = Font(name=_ARIAL, size=10, bold=True)
        c.fill = _FILL_GRIS
        c.alignment = _CENTRO
        c.border = _BORDE
    ws.row_dimensions[4].height = 27

    fila = 5
    for f in facturas:
        valores = [
            f.fecha,
            f.tipo_factura.value if hasattr(f.tipo_factura, "value") else str(f.tipo_factura),
            _num_o_str(f.punto_venta),
            _num_o_str(f.numero_comprobante),
            *[float(getattr(f, col)) or None for col in COLUMNAS_IMPORTE],
            f"=SUM({get_column_letter(_IVA_PRIMERA_COL_IMPORTE)}{fila}:"
            f"{get_column_letter(_IVA_COL_TOTAL - 1)}{fila})",
            f.descripcion,
        ]
        for col, v in enumerate(valores, start=1):
            c = ws.cell(row=fila, column=col, value=v)
            c.font = Font(name=_ARIAL, size=10)
            c.border = _BORDE
            if col == 1:
                c.number_format = "d/mm/yyyy"
                c.alignment = Alignment(horizontal="center")
            elif col in (2, 3, 4):
                c.alignment = Alignment(horizontal="center")
            elif _IVA_PRIMERA_COL_IMPORTE <= col <= _IVA_COL_TOTAL:
                c.number_format = _FMT_IMPORTE
        fila += 1

    # Fila de totales
    ws.cell(row=fila, column=1, value="TOTALES")
    ws.merge_cells(start_row=fila, start_column=1, end_row=fila, end_column=_IVA_PRIMERA_COL_IMPORTE - 1)
    for col in range(1, ultima_col + 1):
        c = ws.cell(row=fila, column=col)
        c.font = Font(name=_ARIAL, size=10, bold=True)
        c.fill = _FILL_GRIS
        c.border = _BORDE
        if col == 1:
            c.alignment = _CENTRO
        if _IVA_PRIMERA_COL_IMPORTE <= col <= _IVA_COL_TOTAL:
            letra = get_column_letter(col)
            c.value = f"=SUM({letra}5:{letra}{fila - 1})" if fila > 5 else 0
            c.number_format = _FMT_IMPORTE

    for col, ancho in enumerate(_IVA_ANCHOS, start=1):
        ws.column_dimensions[get_column_letter(col)].width = ancho
    ws.freeze_panes = "A5"

    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()
