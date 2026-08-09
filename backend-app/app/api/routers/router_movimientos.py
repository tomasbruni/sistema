from fastapi import APIRouter, Depends, HTTPException, status
from fastapi.responses import StreamingResponse
from sqlmodel import Session, SQLModel, select
from sqlalchemy import func
from typing import Optional
from datetime import datetime, date
from io import BytesIO

import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment

from app.db.session import get_session
from app.db.models import MovimientoStock, Accesorio, Local, TipoMovimiento

from app.api.deps import get_current_user, require_admin
from app.api.funciones.fechas import start_of_day, end_of_day

router = APIRouter(
    prefix="/movimientos",
    tags=["MOVIMIENTOS DE STOCK"],
    dependencies=[Depends(require_admin)],
)


# ─── RESPONSE SCHEMA ─────────────────────────────────────────────────────────

class MovimientoResponse(SQLModel):
    id: int
    accesorio_id: int
    accesorio_nombre: Optional[str] = None
    local_id: int
    local_nombre: Optional[str] = None
    tipo_movimiento: TipoMovimiento
    cantidad: int
    stock_anterior: Optional[int] = None
    stock_nuevo: Optional[int] = None
    fecha: Optional[datetime] = None
    motivo: Optional[str] = None
    usuario_id: Optional[int] = None


class MovimientosListResponse(SQLModel):
    total: int
    items: list[MovimientoResponse]


# ─── HELPER: query base compartida con exportación ───────────────────────────

def _build_query(
    local_id, accesorio_id, tipo_movimiento, fecha_desde, fecha_hasta, usuario_id
):
    stmt = (
        select(MovimientoStock, Accesorio, Local)
        .join(Accesorio, MovimientoStock.accesorio_id == Accesorio.accesorio_id)  # type: ignore
        .join(Local, MovimientoStock.local_id == Local.local_id)  # type: ignore
    )
    if local_id is not None:
        stmt = stmt.where(MovimientoStock.local_id == local_id)
    if accesorio_id is not None:
        stmt = stmt.where(MovimientoStock.accesorio_id == accesorio_id)
    if usuario_id is not None:
        stmt = stmt.where(MovimientoStock.usuario_id == usuario_id)
    if tipo_movimiento is not None:
        stmt = stmt.where(MovimientoStock.tipo_movimiento == tipo_movimiento)
    if fecha_desde is not None:
        stmt = stmt.where(MovimientoStock.fecha >= start_of_day(fecha_desde))  # type: ignore
    if fecha_hasta is not None:
        stmt = stmt.where(MovimientoStock.fecha <= end_of_day(fecha_hasta))  # type: ignore
    return stmt.order_by(MovimientoStock.fecha.desc())  # type: ignore


# ─── ENDPOINTS ───────────────────────────────────────────────────────────────

@router.get("/", response_model=MovimientosListResponse)
def listar_movimientos(
    skip: int = 0,
    limit: int = 100,
    # Filtros
    local_id: Optional[int] = None,
    accesorio_id: Optional[int] = None,
    tipo_movimiento: Optional[TipoMovimiento] = None,
    fecha_desde: Optional[date] = None,
    fecha_hasta: Optional[date] = None,
    usuario_id: Optional[int] = None,
    session: Session = Depends(get_session),
):
    """
    Lista el historial de movimientos de stock.

    Filtros disponibles:
    - **local_id**: filtra por local
    - **accesorio_id**: filtra por accesorio
    - **usuario_id**: filtra por usuario que registró el movimiento
    - **tipo_movimiento**: ENTRADA | SALIDA | AJUSTE | VENTA | DEVOLUCION | RESERVA
    - **fecha_desde** / **fecha_hasta**: rango de fechas (formato YYYY-MM-DD, inclusive en ambos extremos)
    """
    base = _build_query(local_id, accesorio_id, tipo_movimiento, fecha_desde, fecha_hasta, usuario_id)

    # Total sin paginar (para que el front pueda mostrar cuántos hay)
    count_stmt = select(func.count()).select_from(base.subquery())
    total = session.exec(count_stmt).one()

    # Resultado paginado (el orden ya viene del helper)
    stmt = base.offset(skip).limit(limit)
    rows = session.exec(stmt).all()

    items = [
        MovimientoResponse(
            id=mov.id,  # type: ignore
            accesorio_id=mov.accesorio_id,
            accesorio_nombre=acc.nombre,
            local_id=mov.local_id,
            local_nombre=loc.nombre,
            tipo_movimiento=mov.tipo_movimiento,
            cantidad=mov.cantidad,
            stock_anterior=mov.stock_anterior,
            stock_nuevo=mov.stock_nuevo,
            fecha=mov.fecha,
            motivo=mov.motivo,
            usuario_id=mov.usuario_id,
        )
        for mov, acc, loc in rows
    ]

    return MovimientosListResponse(total=total, items=items)


@router.get("/export")
def exportar_movimientos(
    local_id: Optional[int] = None,
    accesorio_id: Optional[int] = None,
    tipo_movimiento: Optional[TipoMovimiento] = None,
    fecha_desde: Optional[date] = None,
    fecha_hasta: Optional[date] = None,
    usuario_id: Optional[int] = None,
    session: Session = Depends(get_session),
):
    """Exporta el historial de movimientos filtrado a un archivo Excel."""

    stmt = _build_query(local_id, accesorio_id, tipo_movimiento, fecha_desde, fecha_hasta, usuario_id)
    resultados = session.exec(stmt).all()

    # ── Armar el Excel ────────────────────────────────────────────────────────
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Movimientos"  # type: ignore

    header_font  = Font(bold=True, color="FFFFFF")
    header_fill  = PatternFill(fill_type="solid", fgColor="1A1A2E")
    header_align = Alignment(horizontal="center")

    columnas = ["ID", "Fecha", "Accesorio", "Local", "Tipo", "Cantidad", "Motivo", "Usuario"]
    for col_idx, titulo in enumerate(columnas, start=1):
        cell = ws.cell(row=1, column=col_idx, value=titulo)  # type: ignore
        cell.font  = header_font
        cell.fill  = header_fill
        cell.alignment = header_align

    COLORES_TIPO = {
        TipoMovimiento.ENTRADA: "D4EDDA",
        TipoMovimiento.SALIDA:  "F8D7DA",
        TipoMovimiento.AJUSTE:  "FFF3CD",
        TipoMovimiento.VENTA:   "D1ECF1",
    }

    for mov, acc, loc in resultados:
        fecha_str = (
            mov.fecha.strftime("%d/%m/%Y %H:%M") if mov.fecha else ""
        )
        fila = [
            mov.id,
            fecha_str,
            acc.nombre,
            loc.nombre,
            mov.tipo_movimiento.value,
            mov.cantidad,
            mov.motivo or "",
            mov.usuario_id or "",
        ]
        ws.append(fila)  # type: ignore

        # Colorear la fila según tipo de movimiento
        color = COLORES_TIPO.get(mov.tipo_movimiento)
        if color:
            fill = PatternFill(fill_type="solid", fgColor=color)
            for col_idx in range(1, len(columnas) + 1):
                ws.cell(row=ws.max_row, column=col_idx).fill = fill  # type: ignore

    anchos = {"A": 8, "B": 18, "C": 14, "D": 40, "E": 20, "F": 12, "G": 12, "H": 40, "I": 10}
    for col_letra, ancho in anchos.items():
        ws.column_dimensions[col_letra].width = ancho  # type: ignore

    # ── Respuesta como stream ─────────────────────────────────────────────────
    buffer = BytesIO()
    wb.save(buffer)
    buffer.seek(0)

    nombre_archivo = "movimientos"
    if fecha_desde or fecha_hasta:
        rango = f"_{fecha_desde or ''}_a_{fecha_hasta or ''}"
        nombre_archivo += rango

    return StreamingResponse(
        buffer,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f"attachment; filename={nombre_archivo}.xlsx"},
    )


@router.get("/{movimiento_id}", response_model=MovimientoResponse)
def obtener_movimiento(
    movimiento_id: int,
    session: Session = Depends(get_session),
):
    """Obtiene un movimiento específico por ID."""
    row = session.exec(
        select(MovimientoStock, Accesorio, Local)
        .join(Accesorio, MovimientoStock.accesorio_id == Accesorio.accesorio_id)  # type: ignore
        .join(Local, MovimientoStock.local_id == Local.local_id)  # type: ignore
        .where(MovimientoStock.id == movimiento_id)
    ).first()

    if not row:
        raise HTTPException(
            status_code=status.HTTP_404_NOT_FOUND,
            detail=f"Movimiento con ID {movimiento_id} no encontrado",
        )

    mov, acc, loc = row
    return MovimientoResponse(
        id=mov.id,  # type: ignore
        accesorio_id=mov.accesorio_id,
        accesorio_nombre=acc.nombre,
        local_id=mov.local_id,
        local_nombre=loc.nombre,
        tipo_movimiento=mov.tipo_movimiento,
        cantidad=mov.cantidad,
        stock_anterior=mov.stock_anterior,
        stock_nuevo=mov.stock_nuevo,
        fecha=mov.fecha,
        motivo=mov.motivo,
        usuario_id=mov.usuario_id,
    )

