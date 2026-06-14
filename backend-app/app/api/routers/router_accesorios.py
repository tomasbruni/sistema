from fastapi import APIRouter, Depends, HTTPException, UploadFile, File, Form
from sqlmodel import Session
from pydantic import BaseModel

from typing import Optional, List
from sqlmodel import Session, SQLModel, select, col
from sqlalchemy import exc

from app.db.session import get_session
from app.db.models import *

from app.api.modelscreate import *
from app.api.modelsupdate import *
from app.api.funciones.accesorios_funciones import generar_nombre_accesorio, normalizar_nombre_accesorio
from app.api.funciones.movimientos_stock import aplicar_movimiento_stock
from app.api.deps import get_current_user, require_admin, UsuarioActual

from fastapi.responses import StreamingResponse
from openpyxl import Workbook
from io import BytesIO
from openpyxl.styles import Font, PatternFill, Alignment

import json
import pandas as pd

router = APIRouter(prefix="/accesorios",
                   tags=["ACCESORIOS"],
                   responses={404: {"message": "No encontrado"}})


class SeedAccesorioItem(BaseModel):
    tipo: str
    subtipo: Optional[str] = None
    nombre: str
    precio: int


class SeedAccesoriosRequest(BaseModel):
    data: List[SeedAccesorioItem]


@router.post("/seed")
def seed_accesorios(
    request: SeedAccesoriosRequest,
    session: Session = Depends(get_session),
    current_user: UsuarioActual = Depends(require_admin)
):
    """
    Carga masiva de accesorios referenciando tipo y subtipo por nombre.
    Es idempotente: si ya existe un accesorio con el mismo nombre+tipo+subtipo lo omite.
    Crea entradas de stock en 0 para cada local.
    """
    locales = session.exec(select(Local)).all()
    if not locales:
        raise HTTPException(status_code=400, detail="No hay locales registrados")

    creados = []
    omitidos = []
    errores = []

    for item in request.data:
        # Resolver tipo por nombre
        tipo = session.exec(
            select(TipoAccesorio).where(TipoAccesorio.nombre == item.tipo)
        ).first()
        if not tipo:
            errores.append({"item": item.nombre, "motivo": f"Tipo '{item.tipo}' no encontrado"})
            continue

        # Resolver subtipo por nombre (opcional)
        subtipo_id = None
        if item.subtipo:
            subtipo = session.exec(
                select(SubtipoAccesorio).where(
                    SubtipoAccesorio.tipo_id == tipo.tipo_id,
                    SubtipoAccesorio.nombre == item.subtipo
                )
            ).first()
            if not subtipo:
                errores.append({"item": item.nombre, "motivo": f"Subtipo '{item.subtipo}' no encontrado para tipo '{item.tipo}'"})
                continue
            subtipo_id = subtipo.subtipo_id

        # Verificar duplicado
        stmt = select(Accesorio).where(
            Accesorio.nombre == item.nombre,
            Accesorio.tipo_id == tipo.tipo_id,
            Accesorio.activo == True,  # type: ignore
        )
        stmt = stmt.where(Accesorio.subtipo_id == subtipo_id)
        if session.exec(stmt).first():
            omitidos.append(item.nombre)
            continue

        try:
            accesorio = Accesorio( #type: ignore
                nombre=item.nombre,
                precio=item.precio,
                tipo_id=tipo.tipo_id,
                subtipo_id=subtipo_id,
            )
            session.add(accesorio)
            session.flush()

            session.add_all([
                StockAccesorio(accesorio_id=accesorio.accesorio_id, local_id=local.local_id, cantidad=0)  # type: ignore
                for local in locales
            ])
            creados.append({"nombre": item.nombre, "tipo": item.tipo, "subtipo": item.subtipo})
        except Exception as e:
            session.rollback()
            errores.append({"item": item.nombre, "motivo": str(e)})
            continue

    session.commit()

    return {
        "creados": creados,
        "omitidos": omitidos,
        "errores": errores,
    }


@router.get("/export")
def exportar_accesorios(
    buscar: Optional[str] = "",
    tipo_id: Optional[int] = None,
    subtipo_id: Optional[int] = None,
    activo: Optional[bool] = True,
    session: Session = Depends(get_session)
):
    statement = select(Accesorio)
    if activo is not None:
        statement = statement.where(Accesorio.activo == activo)
    if buscar:
        statement = statement.where(Accesorio.nombre.ilike(f"%{buscar}%"))  # type: ignore
    if tipo_id is not None:
        statement = statement.where(Accesorio.tipo_id == tipo_id)
    if subtipo_id is not None:
        statement = statement.where(Accesorio.subtipo_id == subtipo_id)

    accesorios = session.exec(statement).all()

    wb = Workbook()
    ws = wb.active
    ws.title = "Accesorios"  # type: ignore

    header_font  = Font(bold=True, color="FFFFFF")
    header_fill  = PatternFill(fill_type="solid", fgColor="1A1A2E")
    header_align = Alignment(horizontal="center")

    columnas = ["ID", "Nombre", "Tipo ID", "Subtipo ID", "Marca Celular ID", "Modelo Celular ID", "Precio", "Activo"]
    for col_idx, titulo in enumerate(columnas, start=1):
        cell = ws.cell(row=1, column=col_idx, value=titulo)  # type: ignore
        cell.font      = header_font
        cell.fill      = header_fill
        cell.alignment = header_align

    for a in accesorios:
        ws.append([  # type: ignore
            a.accesorio_id,
            a.nombre,
            a.tipo_id,
            a.subtipo_id,
            a.marca_celular_id,
            a.modelo_celular_id,
            a.precio,
            a.activo,
        ])

    anchos = {"A": 8, "B": 40, "C": 10, "D": 12, "E": 16, "F": 18, "G": 12, "H": 10}
    for col_letra, ancho in anchos.items():
        ws.column_dimensions[col_letra].width = ancho  # type: ignore

    buffer = BytesIO()
    wb.save(buffer)
    buffer.seek(0)

    return StreamingResponse(
        buffer,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": "attachment; filename=accesorios.xlsx"}
    )


# @router.post("/sugerir-nombre")
# def sugerir_nombre_accesorio(
#     data: AccesorioCreate,
#     session: Session = Depends(get_session)
# ):
#     nombre_base = generar_nombre_accesorio(
#         tipo_id=data.tipo_id,
#         session=session,
#         subtipo_id=data.subtipo_id,
#         marca_id=data.marca_id,
#         marca_celular_id=data.marca_celular_id,
#         modelo_celular_id=data.modelo_celular_id
#     )
#     return {"nombre_sugerido": nombre_base}


# NO SIRVE
@router.post("/verificar-duplicado")
def verificar_duplicado(
    data: AccesorioCreate,
    session: Session = Depends(get_session)
):
    statement = select(Accesorio).where(
        Accesorio.nombre  == data.nombre,
        Accesorio.tipo_id == data.tipo_id,  # type: ignore
        Accesorio.activo  == True           # type: ignore
    )

    if data.subtipo_id:
        statement = statement.where(Accesorio.subtipo_id == data.subtipo_id)
    else:
        statement = statement.where(Accesorio.subtipo_id == None)

    if data.marca_celular_id:
        statement = statement.where(Accesorio.marca_celular_id == data.marca_celular_id)
    else:
        statement = statement.where(Accesorio.marca_celular_id == None)

    if data.modelo_celular_id:
        statement = statement.where(Accesorio.modelo_celular_id == data.modelo_celular_id)
    else:
        statement = statement.where(Accesorio.modelo_celular_id == None)

    if data.marca_id:
        statement = statement.where(Accesorio.marca_id == data.marca_id)
    else:
        statement = statement.where(Accesorio.marca_id == None)

    similares = session.exec(statement).all()

    if not similares:
        return {"tiene_duplicados": False, "duplicados": []}

    duplicado_exacto = next((s for s in similares if s.precio == data.precio), None)

    return {
        "tiene_duplicados": True,
        "es_duplicado_exacto": duplicado_exacto is not None,
        "duplicados": [
            {
                "accesorio_id": s.accesorio_id,
                "nombre": s.nombre,
                "precio": s.precio,
                "es_mismo_precio": s.precio == data.precio
            }
            for s in similares
        ]
    }


@router.post("/")
def crear_accesorio(
    accesorio_data: AccesorioCreate,
    session: Session = Depends(get_session),
    current_user: UsuarioActual = Depends(require_admin)
):
    # No se puede especificar marca_celular sola sin modelo, ni modelo sin marca_celular
    if accesorio_data.modelo_celular_id and not accesorio_data.marca_celular_id:
        raise HTTPException(
            status_code=400,
            detail="Si se especifica un modelo de celular, se debe especificar también la marca"
        )

    # Validar que el modelo pertenece a la marca si ambos vienen
    if accesorio_data.marca_celular_id and accesorio_data.modelo_celular_id:
        modelo = session.get(ModeloCelular, accesorio_data.modelo_celular_id)
        if not modelo:
            raise HTTPException(status_code=404, detail="Modelo de celular no encontrado")
        if modelo.marca_celular_id != accesorio_data.marca_celular_id:
            raise HTTPException(status_code=400, detail="El modelo no pertenece a la marca indicada")

    try:
        locales = session.exec(select(Local)).all()
        if not locales:
            raise HTTPException(status_code=400, detail="No hay locales")

        accesorio = Accesorio(**accesorio_data.model_dump())
        accesorio.nombre = normalizar_nombre_accesorio(accesorio.nombre)
        session.add(accesorio)
        session.flush()

        lista_stocks = [
            StockAccesorio(
                accesorio_id=accesorio.accesorio_id,  # type: ignore
                local_id=local.local_id,              # type: ignore
                cantidad=0
            )
            for local in locales
        ]
        session.add_all(lista_stocks)
        session.commit()
        session.refresh(accesorio)
        for stock in lista_stocks:
            session.refresh(stock)

        return {"accesorio": accesorio, "lista-stocks": lista_stocks}

    except ValueError as e:
        session.rollback()
        raise HTTPException(status_code=400, detail=str(e))
    except exc.IntegrityError as e:
        session.rollback()
        raise HTTPException(status_code=400, detail=str(e.orig))
    except Exception as e:
        session.rollback()
        raise HTTPException(status_code=500, detail=f"Error al crear accesorio: {str(e)}")


@router.get("/con-stock")
def listar_accesorios_con_stock(
    buscar: Optional[str] = "",
    local_id: Optional[int] = None,
    session: Session = Depends(get_session),
    current_user: UsuarioActual = Depends(get_current_user)
):
    statement = select(Accesorio).where(Accesorio.activo == True).limit(100)  # type: ignore
    if buscar:
        statement = statement.where(Accesorio.nombre.ilike(f"%{buscar}%"))  # type: ignore

    accesorios = session.exec(statement).all()

    result = []
    for a in accesorios:
        stock = 0
        if local_id is not None:
            stock_row = session.exec(
                select(StockAccesorio).where(
                    StockAccesorio.accesorio_id == a.accesorio_id,
                    StockAccesorio.local_id == local_id,
                )
            ).first()
            stock = stock_row.cantidad if stock_row else 0
        result.append({
            "accesorio_id": a.accesorio_id,
            "nombre": a.nombre,
            "precio": a.precio,
            "stock": stock,
        })
    return result


@router.get("/{accesorio_id}")
def obtener_accesorio(
    accesorio_id: int,
    session: Session = Depends(get_session),
    current_user: UsuarioActual = Depends(get_current_user)
):
    accesorio = session.get(Accesorio, accesorio_id)
    if not accesorio:
        raise HTTPException(status_code=404, detail="Accesorio no encontrado")
    return accesorio


@router.patch("/{accesorio_id}")
def actualizar_accesorio(
    accesorio_id: int,
    accesorio_data: AccesorioUpdate,
    session: Session = Depends(get_session),
    current_user: UsuarioActual = Depends(require_admin)
):
    accesorio = session.get(Accesorio, accesorio_id)
    if not accesorio:
        raise HTTPException(status_code=404, detail="Accesorio no encontrado")

    try:
        update_data = accesorio_data.model_dump(exclude_unset=True)

        if "nombre" in update_data and update_data["nombre"] is not None:
            update_data["nombre"] = normalizar_nombre_accesorio(update_data["nombre"])

        # Validar consistencia marca/modelo si alguno de los dos viene en el update
        marca_celular_id = update_data.get("marca_celular_id", accesorio.marca_celular_id)
        modelo_celular_id = update_data.get("modelo_celular_id", accesorio.modelo_celular_id)

        if modelo_celular_id and not marca_celular_id:
            raise HTTPException(status_code=400, detail="Si se especifica un modelo, se debe especificar también la marca")

        if marca_celular_id and modelo_celular_id:
            modelo = session.get(ModeloCelular, modelo_celular_id)
            if not modelo:
                raise HTTPException(status_code=404, detail="Modelo de celular no encontrado")
            if modelo.marca_celular_id != marca_celular_id:
                raise HTTPException(status_code=400, detail="El modelo no pertenece a la marca indicada")

        for key, value in update_data.items():
            setattr(accesorio, key, value)

        session.add(accesorio)
        session.commit()
        session.refresh(accesorio)
        return accesorio

    except HTTPException:
        session.rollback()
        raise
    except exc.IntegrityError:
        session.rollback()
        raise HTTPException(
            status_code=400,
            detail="No se pudo actualizar el accesorio. Verifique que no haya duplicados."
        )


@router.get("/", response_model=list[Accesorio])
def listar_accesorios(
    skip: int = 0,
    limit: int = 100,
    buscar: Optional[str] = "",
    tipo_id: Optional[int] = None,
    subtipo_id: Optional[int] = None,
    marca_celular_id: Optional[int] = None,
    modelo_celular_id: Optional[int] = None,
    activo: Optional[bool] = True,
    session: Session = Depends(get_session),
    current_user: UsuarioActual = Depends(get_current_user)
):
    statement = select(Accesorio).offset(skip).limit(limit)

    if activo is not None:
        statement = statement.where(Accesorio.activo == activo)
    if buscar:
        statement = statement.where(Accesorio.nombre.ilike(f"%{buscar}%"))  # type: ignore
    if tipo_id is not None:
        statement = statement.where(Accesorio.tipo_id == tipo_id)
    if subtipo_id is not None:
        statement = statement.where(Accesorio.subtipo_id == subtipo_id)
    if marca_celular_id is not None:
        statement = statement.where(Accesorio.marca_celular_id == marca_celular_id)
    if modelo_celular_id is not None:
        statement = statement.where(Accesorio.modelo_celular_id == modelo_celular_id)

    return session.exec(statement).all()


class CambioPrecioMasivoInput(BaseModel):
    tipo_id: int
    subtipo_id: Optional[int] = None
    nuevo_precio: int


@router.post("/cambio-precio-masivo")
def cambio_precio_masivo(
    data: CambioPrecioMasivoInput,
    session: Session = Depends(get_session),
    current_user: UsuarioActual = Depends(require_admin),
):
    if data.nuevo_precio < 0:
        raise HTTPException(status_code=400, detail="El precio no puede ser negativo")

    statement = select(Accesorio).where(Accesorio.tipo_id == data.tipo_id)
    if data.subtipo_id is not None:
        statement = statement.where(Accesorio.subtipo_id == data.subtipo_id)

    accesorios = session.exec(statement).all()
    for acc in accesorios:
        acc.precio = data.nuevo_precio
    session.commit()
    return {"mensaje": f"Precio actualizado a ${data.nuevo_precio:,}", "accesorios_modificados": len(accesorios)}


@router.delete("/{accesorio_id}")
def eliminar_accesorio(
    accesorio_id: int,
    session: Session = Depends(get_session),
    current_user: UsuarioActual = Depends(require_admin)
):
    """
    Elimina físicamente si nunca tuvo ventas ni movimientos.
    Si tiene ventas o movimientos, lo desactiva.
    """
    try:
        accesorio = session.get(Accesorio, accesorio_id)
        if not accesorio:
            raise HTTPException(status_code=404, detail="Accesorio no encontrado")

        tiene_ventas = session.exec(
            select(DetalleVentaAccesorio).where(DetalleVentaAccesorio.accesorio_id == accesorio_id)
        ).first() is not None

        tiene_movimientos = session.exec(
            select(MovimientoStock).where(MovimientoStock.accesorio_id == accesorio_id)
        ).first() is not None

        if tiene_ventas or tiene_movimientos:
            accesorio.activo = False
            session.commit()

            razones = []
            if tiene_ventas:
                razones.append("ventas asociadas")
            if tiene_movimientos:
                razones.append("movimientos de stock")

            return {
                "mensaje": f"Accesorio desactivado ({', '.join(razones)})",
                "accion": "desactivado",
                "accesorio_id": accesorio_id,
                "razones": razones
            }

        else:
            stocks = session.exec(
                select(StockAccesorio).where(StockAccesorio.accesorio_id == accesorio_id)
            ).all()
            for stock in stocks:
                session.delete(stock)
            session.flush()
            session.delete(accesorio)
            session.commit()

            return {
                "mensaje": "Accesorio eliminado correctamente",
                "accion": "eliminado",
                "accesorio_id": accesorio_id
            }

    except HTTPException:
        raise
    except Exception as e:
        session.rollback()
        raise HTTPException(status_code=500, detail=f"Error al eliminar accesorio: {str(e)}")


# ─── IMPORTACIÓN DESDE EXCEL / ODS ───────────────────────────────────────────

@router.post("/importar-excel")
def importar_desde_excel(
    archivo:      UploadFile = File(...),
    tipo_id:      int        = Form(...),
    local_id:     int        = Form(...),
    receptor_id:  int        = Form(...),
    precios:      str        = Form(...),   # JSON: {"SILICONA": 9000, "TRANSPARENTE": 8000, ...}
    observaciones: str       = Form(""),
    session:      Session    = Depends(get_session),
    current_user: UsuarioActual = Depends(require_admin),
):
    """
    Importa accesorios y stock inicial desde un archivo Excel (.xlsx) u ODS (.ods).

    Formato esperado del archivo:
      - Columnas obligatorias: MARCA, MODELO
      - Una columna por subtipo (ej. SILICONA, TRANSPARENTE…) con la cantidad de stock a ingresar

    Parámetros:
      - tipo_id:     ID del tipo de accesorio al que pertenecen todos los ítems del archivo
      - local_id:    local donde se registra el ingreso de stock
      - receptor_id: usuario que recibe la mercadería
      - precios:     JSON con precio por subtipo  {"SILICONA": 9000, ...}
      - observaciones: texto libre para el ingreso de lote (opcional)

    Proceso:
      1. Lee el archivo y normaliza columnas
      2. Crea marcas y modelos que no existan (idempotente)
      3. Crea accesorios que no existan (idempotente)
      4. Si hay cantidades > 0, genera un ingreso de lote

    Devuelve un resumen con contadores y lista de errores por ítem.
    """

    # ── 1. Parsear precios ────────────────────────────────────────────────────
    try:
        precios_dict: dict[str, int] = {
            k.strip().upper(): int(v)
            for k, v in json.loads(precios).items()
        }
    except Exception:
        raise HTTPException(status_code=400, detail="El campo 'precios' debe ser un JSON válido: {\"SILICONA\": 9000, ...}")

    # ── 2. Leer el archivo con pandas ─────────────────────────────────────────
    nombre_archivo = archivo.filename or ""
    contenido      = archivo.file.read()

    try:
        if nombre_archivo.endswith(".ods"):
            df = pd.read_excel(BytesIO(contenido), sheet_name=0, engine="odf")
        elif nombre_archivo.endswith((".xlsx", ".xls")):
            df = pd.read_excel(BytesIO(contenido), sheet_name=0)
        else:
            raise HTTPException(status_code=400, detail="Formato no soportado. Usá .xlsx o .ods")
    except HTTPException:
        raise
    except Exception as e:
        raise HTTPException(status_code=400, detail=f"No se pudo leer el archivo: {e}")

    # Normalizar nombres de columna: sin espacios extra, todo en mayúsculas
    df.columns = [str(c).strip().upper() for c in df.columns]

    if "MARCA" not in df.columns or "MODELO" not in df.columns:
        raise HTTPException(status_code=400, detail="El archivo debe tener columnas MARCA y MODELO")

    # ── 3. Resolver subtipos disponibles para este tipo_id ────────────────────
    # Solo procesamos columnas que coincidan con un subtipo registrado en la DB
    # para el tipo recibido, Y que además tengan un precio definido.
    subtipos_db: dict[str, int] = {
        s.nombre.strip().upper(): s.subtipo_id   # type: ignore
        for s in session.exec(
            select(SubtipoAccesorio).where(SubtipoAccesorio.tipo_id == tipo_id)
        ).all()
    }

    # Columnas del archivo que son subtipos reconocidos con precio definido
    subtipos_activos: dict[str, int] = {
        nombre: subtipo_id
        for nombre, subtipo_id in subtipos_db.items()
        if nombre in df.columns and nombre in precios_dict
    }

    if not subtipos_activos:
        raise HTTPException(
            status_code=400,
            detail=(
                "Ninguna columna del archivo coincide con subtipos registrados para ese tipo_id "
                "o no tienen precio definido. "
                f"Subtipos en DB: {list(subtipos_db.keys())}. "
                f"Columnas en archivo: {list(df.columns)}."
            )
        )

    # Advertencias sobre columnas ignoradas:
    # subtipos con columna en el archivo pero sin precio definido
    advertencias: list[str] = []
    for nombre in subtipos_db:
        if nombre in df.columns and nombre not in precios_dict:
            advertencias.append(f"Subtipo '{nombre}' está en el archivo pero no tiene precio definido — se ignoró")
    # precios definidos que no coinciden con ningún subtipo de la DB
    for nombre in precios_dict:
        if nombre not in subtipos_db:
            advertencias.append(f"Precio definido para '{nombre}' pero no es un subtipo registrado para este tipo — se ignoró")
    # precio definido y subtipo en DB, pero sin columna en el archivo
    for nombre in precios_dict:
        if nombre in subtipos_db and nombre not in df.columns:
            advertencias.append(f"Precio definido para '{nombre}' pero no hay columna '{nombre}' en el archivo — se ignoró")

    # ── 4. Verificar que el tipo y el local existan ───────────────────────────
    if not session.get(TipoAccesorio, tipo_id):
        raise HTTPException(status_code=404, detail=f"Tipo de accesorio {tipo_id} no encontrado")

    locales = session.exec(select(Local)).all()
    if not any(l.local_id == local_id for l in locales):
        raise HTTPException(status_code=404, detail=f"Local {local_id} no encontrado")

    # ── 5. Crear marcas y modelos ─────────────────────────────────────────────
    # Acumulamos IDs en memoria para no hacer una query por cada fila del loop
    marcas_creadas    = 0
    marcas_existentes = 0
    modelos_creados    = 0
    modelos_existentes = 0

    # marca_nombre → marca_celular_id
    marcas_ids: dict[str, int] = {}
    # (marca_nombre, modelo_nombre) → modelo_celular_id
    modelos_ids: dict[tuple[str, str], int] = {}

    errores: list[dict] = []
    filas_invalidas: set[int] = set()  # índices de filas a saltear en el loop de accesorios

    for idx_raw, row in df.iterrows():
        idx = int(idx_raw)  # type: ignore[arg-type]
        marca_raw  = row["MARCA"]
        modelo_raw = row["MODELO"]

        # Saltear filas con MARCA o MODELO vacío o NaN
        fila_num = int(idx) + 2  # +2 porque el Excel empieza en 1 y la fila 1 es el header
        if pd.isna(marca_raw) or str(marca_raw).strip() == "":
            errores.append({"fila": f"Fila {fila_num}", "motivo": "MARCA vacía o inválida"})
            filas_invalidas.add(idx)
            continue
        if pd.isna(modelo_raw) or str(modelo_raw).strip() == "":
            errores.append({"fila": f"Fila {fila_num}", "motivo": "MODELO vacío o inválido"})
            filas_invalidas.add(idx)
            continue

        marca_nombre  = str(marca_raw).strip().upper()
        modelo_nombre = str(modelo_raw).strip().upper()

        # Crear la marca si no existe, o reusar la existente
        if marca_nombre not in marcas_ids:
            marca = session.exec(
                select(MarcaCelular).where(MarcaCelular.nombre == marca_nombre)
            ).first()

            if marca:
                marcas_existentes += 1
            else:
                marca = MarcaCelular(nombre=marca_nombre)
                session.add(marca)
                session.flush()   # necesitamos el ID antes de seguir
                marcas_creadas += 1

            marcas_ids[marca_nombre] = marca.marca_celular_id  # type: ignore

        marca_id = marcas_ids[marca_nombre]

        # Crear el modelo si no existe bajo esa marca
        if (marca_nombre, modelo_nombre) not in modelos_ids:
            modelo = session.exec(
                select(ModeloCelular).where(
                    ModeloCelular.marca_celular_id == marca_id,
                    ModeloCelular.nombre           == modelo_nombre,
                )
            ).first()

            if modelo:
                modelos_existentes += 1
            else:
                modelo = ModeloCelular(nombre=modelo_nombre, marca_celular_id=marca_id)  # type: ignore
                session.add(modelo)
                session.flush()
                modelos_creados += 1

            modelos_ids[(marca_nombre, modelo_nombre)] = modelo.modelo_celular_id  # type: ignore

    # ── 6. Crear accesorios y acumular ítems para el ingreso ──────────────────
    accesorios_creados  = 0
    accesorios_omitidos = 0

    # Lista de {accesorio_id, cantidad_ingreso} para el ingreso de lote
    items_ingreso: list[dict] = []

    filas_validas = len(df) - len(filas_invalidas)
    combinaciones_esperadas = filas_validas * len(subtipos_activos)

    for idx, row in df.iterrows():
        # Saltear filas que ya se marcaron como inválidas en el paso anterior
        if idx in filas_invalidas:
            continue

        marca_nombre  = str(row["MARCA"]).strip().upper()
        modelo_nombre = str(row["MODELO"]).strip().upper()

        marca_id  = marcas_ids.get(marca_nombre)
        modelo_id = modelos_ids.get((marca_nombre, modelo_nombre))

        if not marca_id or not modelo_id:
            # No debería ocurrir, pero lo capturamos igual
            errores.append({"fila": f"{marca_nombre} / {modelo_nombre}", "motivo": "IDs no resueltos"})
            continue

        for subtipo_nombre, subtipo_id in subtipos_activos.items():
            precio   = precios_dict[subtipo_nombre]
            raw_cantidad = row.get(subtipo_nombre)
            cantidad = 0 if (raw_cantidad is None or pd.isna(raw_cantidad)) else int(raw_cantidad)

            # Nombre descriptivo del accesorio
            nombre_acc = normalizar_nombre_accesorio(f"{subtipo_nombre} {marca_nombre} {modelo_nombre}")

            # Verificar si ya existe un accesorio con esa combinación exacta
            existente = session.exec(
                select(Accesorio).where(
                    Accesorio.nombre            == nombre_acc,
                    Accesorio.tipo_id           == tipo_id,       # type: ignore
                    Accesorio.subtipo_id        == subtipo_id,    # type: ignore
                    Accesorio.marca_celular_id  == marca_id,      # type: ignore ACA SE CLAVO PQ NO DETECTO A13 SIN MARCA Y MODELO
                    Accesorio.modelo_celular_id == modelo_id,     # type: ignore
                    Accesorio.activo            == True,          # type: ignore
                )
            ).first()

            if existente:
                accesorios_omitidos += 1
                # Aunque el accesorio exista, puede tener stock a ingresar.
                # Solo lo agregamos si el stock actual en ese local es 0,
                # para no duplicar si la importación se corre más de una vez.
                if cantidad > 0:
                    stock_actual = session.exec(
                        select(StockAccesorio).where(
                            StockAccesorio.accesorio_id == existente.accesorio_id,
                            StockAccesorio.local_id     == local_id,
                        )
                    ).first()
                    if stock_actual and stock_actual.cantidad == 0:
                        items_ingreso.append({
                            "accesorio_id":    existente.accesorio_id,
                            "cantidad_ingreso": cantidad,
                        })
                continue

            try:
                accesorio = Accesorio(  # type: ignore
                    nombre            = nombre_acc,
                    precio            = precio,
                    tipo_id           = tipo_id,
                    subtipo_id        = subtipo_id,
                    marca_celular_id  = marca_id,
                    modelo_celular_id = modelo_id,
                )
                session.add(accesorio)
                session.flush()  # obtenemos el ID antes de crear los stocks

                # Crear entrada de stock en 0 para cada local
                session.add_all([
                    StockAccesorio(
                        accesorio_id = accesorio.accesorio_id,  # type: ignore
                        local_id     = local.local_id,          # type: ignore
                        cantidad     = 0,
                    )
                    for local in locales
                ])

                accesorios_creados += 1

                # Agregar al ingreso solo si hay cantidad a cargar
                if cantidad > 0:
                    items_ingreso.append({
                        "accesorio_id":    accesorio.accesorio_id,
                        "cantidad_ingreso": cantidad,
                    })

            except Exception as e:
                session.rollback()
                errores.append({"fila": nombre_acc, "motivo": str(e)})
                continue

    session.commit()

    # ── 7. Ingreso de lote ────────────────────────────────────────────────────
    ingreso_lote_id  = None
    unidades_totales = 0

    if items_ingreso:
        lote = IngresoLote(  # type: ignore
            local_id     = local_id,
            receptor_id  = receptor_id,
            observaciones = observaciones or None,
        )
        session.add(lote)
        session.flush()

        for item in items_ingreso:
            stock = session.exec(
                select(StockAccesorio).where(
                    StockAccesorio.accesorio_id == item["accesorio_id"],
                    StockAccesorio.local_id     == local_id,
                )
            ).first()

            if not stock:
                # No debería pasar porque lo creamos arriba, pero por las dudas
                errores.append({"fila": f"accesorio_id={item['accesorio_id']}", "motivo": "Stock no encontrado para ingreso"})
                continue

            # Aplica el delta y graba el movimiento ENTRADA con snapshots
            # permitir_negativo: el ingreso suma stock; si el actual estaba negativo
            # (ventas sin stock), un ingreso parcial deja un residual negativo sin
            # errorear, como flag para recontar.
            aplicar_movimiento_stock(
                session, stock,
                tipo_movimiento = TipoMovimiento.ENTRADA,
                cantidad        = item["cantidad_ingreso"],
                motivo          = f"Importación Excel — {observaciones}" if observaciones else "Importación Excel",
                usuario_id      = current_user.usuario_id,
                ingreso_lote_id = lote.ingreso_lote_id,
                permitir_negativo = True,
            )

            unidades_totales += item["cantidad_ingreso"]

        session.commit()
        ingreso_lote_id = lote.ingreso_lote_id

    # ── 8. Respuesta ──────────────────────────────────────────────────────────
    return {
        "marcas_creadas":           marcas_creadas,
        "marcas_existentes":        marcas_existentes,
        "modelos_creados":          modelos_creados,
        "modelos_existentes":       modelos_existentes,
        "accesorios_creados":       accesorios_creados,
        "accesorios_omitidos":      accesorios_omitidos,
        "combinaciones_esperadas":  combinaciones_esperadas,
        "errores":                  errores,
        "advertencias":             advertencias,
        "ingreso_lote_id":          ingreso_lote_id,
        "items_ingresados":         len(items_ingreso),
        "unidades_totales":         unidades_totales,
    }